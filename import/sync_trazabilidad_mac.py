"""
Sincroniza la planilla BASE MAC ADDRESS hacia tablas propias de trazabilidad.

Solo estas pestañas:
    Ingreso EBD Oficina
    Egreso EBD 2.0
    Ingreso D3 Oficina
    Egreso D3 2.0 (venta/comodato desde Comercial)

El resto de la planilla no se importa.
"""

import os
import re
from datetime import date

from openpyxl import load_workbook

from importar_stock_cantidad import a_fecha, limpiar_texto
from sync_planilla_drive import conectar, descargar, leer_env, log

SHEET_MAC_ID = "1kzUJvelU8QDqThQkYHE2v2Anc0xW9avSTWQeCGWj_gY"

HOJAS = (
    ("Ingreso EBD Oficina", "ingreso", "trazabilidad_ingreso_ebd"),
    ("Egreso EBD 2.0", "egreso", "trazabilidad_egreso_ebd"),
    ("Ingreso D3 Oficina", "ingreso", "trazabilidad_ingreso_d3"),
    ("Egreso D3 2.0 (venta", "egreso", "trazabilidad_egreso_d3"),
)

ALIASES = {
    "FECHA DE LLEGADA": "fecha_llegada",
    "FECHA DE CARGA": "fecha_carga",
    "MAC WIFI": "mac_wifi",
    "MAC ETH": "mac_eth",
    "LOTE": "lote",
    "CAJA N": "caja",
    "CAJA": "caja",
    "CLIENTE": "cliente",
    "FECHA DE DESPACHO": "fecha_despacho",
}


def sheet_id():
    return leer_env().get("GOOGLE_SHEET_MAC_ID", "").strip() or SHEET_MAC_ID


def norm_header(valor):
    texto = (limpiar_texto(valor) or "").upper()
    texto = texto.replace("°", "").replace("º", "").replace("ª", "")
    return re.sub(r"\s+", " ", texto).strip()


def celda(vals, indice):
    if indice is None or indice >= len(vals):
        return None
    return vals[indice]


def mac(valor):
    texto = limpiar_texto(valor)
    if not texto:
        return None
    limpio = re.sub(r"[^0-9A-Fa-f]", "", texto).upper()
    if len(limpio) < 8:
        return None
    return limpio[:20]


def texto_corto(valor, largo):
    texto = limpiar_texto(valor)
    if not texto or texto == "-":
        return None
    return texto[:largo]


def caja(valor):
    if isinstance(valor, float) and valor == int(valor):
        return str(int(valor))
    if isinstance(valor, int) and not isinstance(valor, bool):
        return str(valor)
    return texto_corto(valor, 40)


def fecha_o_texto(valor, largo=40):
    iso = a_fecha(valor)
    if iso:
        return iso, None
    texto = limpiar_texto(valor)
    if not texto or texto == "-":
        return None, None
    coincidencia = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", texto)
    if coincidencia:
        dia, mes, anio = (int(coincidencia.group(1)), int(coincidencia.group(2)), int(coincidencia.group(3)))
        try:
            return date(anio, mes, dia).isoformat(), None
        except ValueError:
            return None, texto[:largo]
    return None, texto[:largo]


def clasificar(nombre):
    limpio = nombre.strip()
    for prefijo, tipo, tabla in HOJAS:
        if limpio == prefijo or limpio.startswith(prefijo):
            if prefijo.startswith("Egreso D3 2.0") and not limpio.startswith("Egreso D3 2.0 (venta"):
                continue
            return tipo, tabla
    return None


def fila_vacia(registro):
    return not any(registro.get(clave) for clave in (
        "mac_wifi", "mac_eth", "cliente", "lote", "caja",
        "fecha_llegada", "fecha_carga", "fecha_despacho",
        "fecha_llegada_texto", "fecha_carga_texto", "fecha_despacho_texto",
    ))


def parsear(ruta):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hojas = []
    try:
        for nombre in libro.sheetnames:
            clasificacion = clasificar(nombre)
            if not clasificacion:
                continue
            tipo, tabla = clasificacion
            hoja = libro[nombre]
            columnas = {}
            filas = []
            for numero, row in enumerate(hoja.iter_rows(values_only=True), start=1):
                vals = list(row)
                if any(norm_header(v) == "MAC WIFI" for v in vals):
                    columnas = {}
                    for indice, valor in enumerate(vals):
                        clave = ALIASES.get(norm_header(valor))
                        if clave and clave not in columnas:
                            columnas[clave] = indice
                    continue
                if "mac_wifi" not in columnas:
                    continue
                wifi = mac(celda(vals, columnas.get("mac_wifi")))
                if not wifi:
                    continue
                if tipo == "ingreso":
                    llegada, llegada_texto = fecha_o_texto(celda(vals, columnas.get("fecha_llegada")))
                    carga, carga_texto = fecha_o_texto(celda(vals, columnas.get("fecha_carga")))
                    registro = {
                        "fecha_llegada": llegada,
                        "fecha_llegada_texto": llegada_texto,
                        "fecha_carga": carga,
                        "fecha_carga_texto": carga_texto,
                        "mac_wifi": wifi,
                        "mac_eth": mac(celda(vals, columnas.get("mac_eth"))),
                        "lote": texto_corto(celda(vals, columnas.get("lote")), 80),
                        "caja": caja(celda(vals, columnas.get("caja"))),
                        "origen_fila": numero,
                    }
                else:
                    despacho, despacho_texto = fecha_o_texto(celda(vals, columnas.get("fecha_despacho")), 80)
                    registro = {
                        "cliente": texto_corto(celda(vals, columnas.get("cliente")), 255),
                        "fecha_despacho": despacho,
                        "fecha_despacho_texto": despacho_texto,
                        "mac_wifi": wifi,
                        "mac_eth": mac(celda(vals, columnas.get("mac_eth"))),
                        "lote": texto_corto(celda(vals, columnas.get("lote")), 80),
                        "caja": caja(celda(vals, columnas.get("caja"))),
                        "origen_fila": numero,
                    }
                if fila_vacia(registro):
                    continue
                filas.append(registro)
            hojas.append({"nombre": nombre, "tipo": tipo, "tabla": tabla, "filas": filas})
    finally:
        libro.close()
    if len(hojas) != 4:
        raise RuntimeError(f"La planilla de MAC no trajo las 4 pestañas (llegaron {len(hojas)})")
    return hojas


def asegurar_tablas(cnx):
    ruta = os.path.join(os.path.dirname(__file__), "..", "sql", "migracion_trazabilidad_mac.sql")
    with open(ruta, encoding="utf-8") as archivo:
        script = archivo.read()
    cur = cnx.cursor()
    for sentencia in script.split(";"):
        sentencia = sentencia.strip()
        if sentencia and not sentencia.upper().startswith("USE "):
            cur.execute(sentencia)
    cnx.commit()


def aplicar(cnx, hojas):
    cur = cnx.cursor()
    total = 0
    for hoja in hojas:
        tabla = hoja["tabla"]
        if hoja["tipo"] == "ingreso":
            sql = f"""INSERT INTO {tabla}
                (fecha_llegada, fecha_llegada_texto, fecha_carga, fecha_carga_texto,
                 mac_wifi, mac_eth, lote, caja, origen_fila)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                  fecha_llegada=VALUES(fecha_llegada),
                  fecha_llegada_texto=VALUES(fecha_llegada_texto),
                  fecha_carga=VALUES(fecha_carga),
                  fecha_carga_texto=VALUES(fecha_carga_texto),
                  mac_wifi=VALUES(mac_wifi),
                  mac_eth=VALUES(mac_eth),
                  lote=VALUES(lote),
                  caja=VALUES(caja)"""
            datos = [
                (f["fecha_llegada"], f["fecha_llegada_texto"], f["fecha_carga"], f["fecha_carga_texto"],
                 f["mac_wifi"], f["mac_eth"], f["lote"], f["caja"], f["origen_fila"])
                for f in hoja["filas"]
            ]
        else:
            sql = f"""INSERT INTO {tabla}
                (cliente, fecha_despacho, fecha_despacho_texto,
                 mac_wifi, mac_eth, lote, caja, origen_fila)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                  cliente=VALUES(cliente),
                  fecha_despacho=VALUES(fecha_despacho),
                  fecha_despacho_texto=VALUES(fecha_despacho_texto),
                  mac_wifi=VALUES(mac_wifi),
                  mac_eth=VALUES(mac_eth),
                  lote=VALUES(lote),
                  caja=VALUES(caja)"""
            datos = [
                (f["cliente"], f["fecha_despacho"], f["fecha_despacho_texto"],
                 f["mac_wifi"], f["mac_eth"], f["lote"], f["caja"], f["origen_fila"])
                for f in hoja["filas"]
            ]
        if not datos:
            log.warning("MAC %s no trajo equipos; no se borra la tabla", hoja["nombre"])
            continue
        cur.executemany(sql, datos)
        cur.execute(
            """CREATE TEMPORARY TABLE IF NOT EXISTS tmp_mac_sync (
                 origen_fila INT NOT NULL PRIMARY KEY
               )"""
        )
        cur.execute("DELETE FROM tmp_mac_sync")
        cur.executemany(
            "INSERT IGNORE INTO tmp_mac_sync (origen_fila) VALUES (%s)",
            [(f["origen_fila"],) for f in hoja["filas"]],
        )
        cur.execute(
            f"""DELETE t FROM {tabla} t
                LEFT JOIN tmp_mac_sync s ON s.origen_fila = t.origen_fila
                WHERE s.origen_fila IS NULL"""
        )
        total += len(hoja["filas"])
        log.info("MAC %s: %s equipos", hoja["nombre"], len(hoja["filas"]))
    cur.execute(
        """INSERT INTO sync_control (archivo, ultima_sync, filas_ok, filas_error, detalle)
           VALUES ('MAC_TRAZABILIDAD', NOW(), %s, 0, 'Google Sheets')
           ON DUPLICATE KEY UPDATE
             ultima_sync = NOW(),
             filas_ok = VALUES(filas_ok),
             filas_error = 0,
             detalle = 'Google Sheets'""",
        (total,),
    )
    cnx.commit()
    log.info("Trazabilidad de MAC actualizada: %s equipos en %s pestañas", total, len(hojas))


def aplicar_archivo(ruta):
    hojas = parsear(ruta)
    cnx = conectar()
    try:
        asegurar_tablas(cnx)
        aplicar(cnx, hojas)
    except Exception:
        cnx.rollback()
        raise
    finally:
        cnx.close()


def sincronizar():
    destino = os.path.join(os.environ.get("TEMP", os.environ.get("TMP", "/tmp")), "stock-mac.xlsx")
    descargar(sheet_id(), destino)
    aplicar_archivo(destino)
