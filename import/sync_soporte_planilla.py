"""
Sincroniza IngresosVarios3 y EgresosVarios3 de la planilla de soporte.

Las filas van a ingreso_soporte y egreso_soporte. No se copian a
movimientos: esos equipos ya pueden tener un movimiento de soporte y
repetirlos duplicaría el historial.
"""

import os
import re
from datetime import date, datetime

from openpyxl import load_workbook

from importar_stock_cantidad import a_fecha, limpiar_texto
from sync_planilla_drive import conectar, descargar, leer_env, log

SHEET_SOPORTE_ID = "1b6zpjoQjgn6bkx6zCBZKgWDAjZ0cWG0x4OPEXJVvNHQ"

HOJAS = {
    "IngresosVarios3": "ingreso",
    "EgresosVarios3": "egreso",
}


def sheet_id():
    return leer_env().get("GOOGLE_SHEET_SOPORTE_ID", "").strip() or SHEET_SOPORTE_ID


def norm(valor):
    texto = (limpiar_texto(valor) or "").lower()
    texto = texto.replace("°", "").replace("º", "")
    return re.sub(r"\s+", " ", texto).strip()


def como_texto(valor, largo):
    if valor is None or isinstance(valor, datetime):
        return None
    if isinstance(valor, float) and valor == int(valor):
        valor = str(int(valor))
    elif isinstance(valor, int) and not isinstance(valor, bool):
        valor = str(valor)
    texto = limpiar_texto(valor)
    if not texto or texto == "-":
        return None
    return texto[:largo]


def mac(valor):
    texto = limpiar_texto(valor)
    if not texto:
        return None
    limpio = re.sub(r"[^0-9A-Fa-f]", "", texto).upper()
    if len(limpio) < 8:
        return None
    return limpio[:20]


def fecha_o_texto(valor):
    iso = a_fecha(valor)
    if iso:
        return iso, None
    texto = limpiar_texto(valor)
    if not texto or texto == "-":
        return None, None
    coincidencia = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", texto)
    if coincidencia:
        dia, mes, anio = (int(coincidencia.group(1)), int(coincidencia.group(2)), int(coincidencia.group(3)))
        try:
            return date(anio, mes, dia).isoformat(), None
        except ValueError:
            return None, texto[:40]
    return None, texto[:40]


def indice(header, *nombres):
    buscados = {norm(n) for n in nombres}
    for i, valor in enumerate(header):
        if norm(valor) in buscados:
            return i
    return None


def celda(row, i):
    if i is None or i >= len(row):
        return None
    return row[i]


def parsear(ruta):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hojas = []
    try:
        for nombre, tipo in HOJAS.items():
            if nombre not in libro.sheetnames:
                raise RuntimeError(f"Falta la pestaña {nombre}")
            hoja = libro[nombre]
            filas_iter = hoja.iter_rows(values_only=True)
            header = list(next(filas_iter))
            if tipo == "ingreso":
                cols = {
                    "fecha": indice(header, "fecha_ingreso"),
                    "razsocial": indice(header, "razsocial_agreg"),
                    "razon": indice(header, "razon_social"),
                    "dispositivo": indice(header, "dispositivo_ingreso"),
                    "estado": indice(header, "estado_ingreso"),
                    "mac": indice(header, "mac"),
                    "segui": indice(header, "segui_ingreso"),
                    "correo": indice(header, "correo_ingreso"),
                    "ticket": indice(header, "num_Ticket", "num_ticket"),
                    "obs": indice(header, "obs_encomienda_ingreso"),
                    "contacto": indice(header, "con_soporte_ingreso"),
                    "motivo": indice(header, "motivo"),
                    "instancia": indice(header, "instancia_gestion"),
                    "wifi": indice(header, "MAC WIFI"),
                    "eth": indice(header, "MAC ETH"),
                }
            else:
                cols = {
                    "fecha": indice(header, "fecha_egreso"),
                    "razsocial": indice(header, "razsocial_agreg"),
                    "razon": indice(header, "razon_social"),
                    "dispositivo": indice(header, "dispositivo_egreso"),
                    "estado": indice(header, "estado_egreso"),
                    "mac": indice(header, "mac"),
                    "envio": indice(header, "envio_retiro"),
                    "correo": indice(header, "correo_egreso"),
                    "segui": indice(header, "segui_egreso"),
                    "ticket": indice(header, "num_Ticket", "num_ticket"),
                    "obs": indice(header, "obs_encomienda_egreso"),
                    "contacto": indice(header, "con_soporte_egreso"),
                    "instancia": indice(header, "instancia_gestion"),
                    "wifi": indice(header, "MAC WIFI"),
                    "eth": indice(header, "MAC ETH"),
                }
            filas = []
            for numero, row in enumerate(filas_iter, start=2):
                valores = list(row)
                if any(norm(v) in ("mac wifi", "fecha_ingreso", "fecha_egreso") for v in valores):
                    continue
                fecha, fecha_texto = fecha_o_texto(celda(valores, cols["fecha"]))
                registro = {
                    "fecha": fecha,
                    "fecha_texto": fecha_texto,
                    "razsocial": como_texto(celda(valores, cols["razsocial"]), 255),
                    "razon": como_texto(celda(valores, cols["razon"]), 255),
                    "dispositivo": como_texto(celda(valores, cols["dispositivo"]), 120),
                    "estado": como_texto(celda(valores, cols["estado"]), 80),
                    "mac": como_texto(celda(valores, cols["mac"]), 80),
                    "envio": como_texto(celda(valores, cols.get("envio")), 80) if tipo == "egreso" else None,
                    "correo": como_texto(celda(valores, cols["correo"]), 80),
                    "segui": como_texto(celda(valores, cols["segui"]), 80),
                    "ticket": como_texto(celda(valores, cols["ticket"]), 80),
                    "obs": como_texto(celda(valores, cols["obs"]), 4000),
                    "contacto": como_texto(celda(valores, cols["contacto"]), 255),
                    "motivo": como_texto(celda(valores, cols.get("motivo")), 255) if tipo == "ingreso" else None,
                    "instancia": como_texto(celda(valores, cols["instancia"]), 80),
                    "mac_wifi": mac(celda(valores, cols["wifi"])),
                    "mac_eth": mac(celda(valores, cols["eth"])),
                    "origen_fila": numero,
                }
                if not any(registro[k] for k in (
                    "razsocial", "razon", "dispositivo", "mac", "mac_wifi", "mac_eth", "segui", "ticket", "obs"
                )):
                    continue
                filas.append(registro)
            hojas.append({"nombre": nombre, "tipo": tipo, "filas": filas})
    finally:
        libro.close()
    return hojas


def asegurar_tablas(cnx):
    ruta = os.path.join(os.path.dirname(__file__), "..", "sql", "migracion_soporte_planilla.sql")
    with open(ruta, encoding="utf-8") as archivo:
        script = archivo.read()
    cur = cnx.cursor()
    for sentencia in script.split(";"):
        sentencia = sentencia.strip()
        if sentencia:
            cur.execute(sentencia)
    extras = {
        "ingreso_soporte": (
            ("perifericos", "VARCHAR(255) NULL"),
            ("codigo_seguimiento", "VARCHAR(80) NULL"),
            ("deposito", "VARCHAR(120) NULL"),
            ("lote", "VARCHAR(80) NULL"),
            ("remito", "VARCHAR(80) NULL"),
            ("factura", "VARCHAR(80) NULL"),
            ("caja", "VARCHAR(40) NULL"),
            ("origen", "VARCHAR(20) NOT NULL DEFAULT 'planilla'"),
        ),
        "egreso_soporte": (
            ("origen", "VARCHAR(20) NOT NULL DEFAULT 'planilla'"),
        ),
    }
    for tabla, columnas in extras.items():
        for nombre, tipo in columnas:
            cur.execute(
                """SELECT COUNT(*) FROM information_schema.COLUMNS
                   WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s AND COLUMN_NAME = %s""",
                (tabla, nombre),
            )
            if cur.fetchone()[0] == 0:
                cur.execute(f"ALTER TABLE {tabla} ADD COLUMN {nombre} {tipo}")
    cnx.commit()


def aplicar(cnx, hojas):
    cur = cnx.cursor()
    total = 0
    for hoja in hojas:
        tabla = "ingreso_soporte" if hoja["tipo"] == "ingreso" else "egreso_soporte"
        if hoja["tipo"] == "ingreso":
            sql = f"""INSERT INTO {tabla}
                (fecha_ingreso, fecha_ingreso_texto, razsocial_agreg, razon_social,
                 dispositivo_ingreso, estado_ingreso, mac, segui_ingreso, correo_ingreso,
                 num_ticket, obs_encomienda_ingreso, con_soporte_ingreso, motivo,
                 instancia_gestion, mac_wifi, mac_eth, origen_fila)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                  fecha_ingreso=VALUES(fecha_ingreso),
                  fecha_ingreso_texto=VALUES(fecha_ingreso_texto),
                  razsocial_agreg=VALUES(razsocial_agreg),
                  razon_social=VALUES(razon_social),
                  dispositivo_ingreso=VALUES(dispositivo_ingreso),
                  estado_ingreso=VALUES(estado_ingreso),
                  mac=VALUES(mac),
                  segui_ingreso=VALUES(segui_ingreso),
                  correo_ingreso=VALUES(correo_ingreso),
                  num_ticket=VALUES(num_ticket),
                  obs_encomienda_ingreso=VALUES(obs_encomienda_ingreso),
                  con_soporte_ingreso=VALUES(con_soporte_ingreso),
                  motivo=VALUES(motivo),
                  instancia_gestion=VALUES(instancia_gestion),
                  mac_wifi=VALUES(mac_wifi),
                  mac_eth=VALUES(mac_eth)"""
            datos = [
                (f["fecha"], f["fecha_texto"], f["razsocial"], f["razon"], f["dispositivo"], f["estado"],
                 f["mac"], f["segui"], f["correo"], f["ticket"], f["obs"], f["contacto"], f["motivo"],
                 f["instancia"], f["mac_wifi"], f["mac_eth"], f["origen_fila"])
                for f in hoja["filas"]
            ]
        else:
            sql = f"""INSERT INTO {tabla}
                (fecha_egreso, fecha_egreso_texto, razsocial_agreg, razon_social,
                 dispositivo_egreso, estado_egreso, mac, envio_retiro, correo_egreso,
                 segui_egreso, num_ticket, obs_encomienda_egreso, con_soporte_egreso,
                 instancia_gestion, mac_wifi, mac_eth, origen_fila)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                  fecha_egreso=VALUES(fecha_egreso),
                  fecha_egreso_texto=VALUES(fecha_egreso_texto),
                  razsocial_agreg=VALUES(razsocial_agreg),
                  razon_social=VALUES(razon_social),
                  dispositivo_egreso=VALUES(dispositivo_egreso),
                  estado_egreso=VALUES(estado_egreso),
                  mac=VALUES(mac),
                  envio_retiro=VALUES(envio_retiro),
                  correo_egreso=VALUES(correo_egreso),
                  segui_egreso=VALUES(segui_egreso),
                  num_ticket=VALUES(num_ticket),
                  obs_encomienda_egreso=VALUES(obs_encomienda_egreso),
                  con_soporte_egreso=VALUES(con_soporte_egreso),
                  instancia_gestion=VALUES(instancia_gestion),
                  mac_wifi=VALUES(mac_wifi),
                  mac_eth=VALUES(mac_eth)"""
            datos = [
                (f["fecha"], f["fecha_texto"], f["razsocial"], f["razon"], f["dispositivo"], f["estado"],
                 f["mac"], f["envio"], f["correo"], f["segui"], f["ticket"], f["obs"], f["contacto"],
                 f["instancia"], f["mac_wifi"], f["mac_eth"], f["origen_fila"])
                for f in hoja["filas"]
            ]
        if not datos:
            log.warning("Soporte %s no trajo filas; no se borra la tabla", hoja["nombre"])
            continue
        cur.executemany(sql, datos)
        cur.execute(
            """CREATE TEMPORARY TABLE IF NOT EXISTS tmp_soporte_sync (
                 origen_fila INT NOT NULL PRIMARY KEY
               )"""
        )
        cur.execute("DELETE FROM tmp_soporte_sync")
        cur.executemany(
            "INSERT IGNORE INTO tmp_soporte_sync (origen_fila) VALUES (%s)",
            [(f["origen_fila"],) for f in hoja["filas"]],
        )
        cur.execute(
            f"""DELETE t FROM {tabla} t
                LEFT JOIN tmp_soporte_sync s ON s.origen_fila = t.origen_fila
                WHERE s.origen_fila IS NULL
                  AND COALESCE(t.origen, 'planilla') = 'planilla'"""
        )
        total += len(hoja["filas"])
        log.info("Soporte %s: %s filas", hoja["nombre"], len(hoja["filas"]))
    cur.execute(
        """INSERT INTO sync_control (archivo, ultima_sync, filas_ok, filas_error, detalle)
           VALUES ('SOPORTE_PLANILLA', NOW(), %s, 0, 'Google Sheets')
           ON DUPLICATE KEY UPDATE
             ultima_sync = NOW(), filas_ok = VALUES(filas_ok),
             filas_error = 0, detalle = 'Google Sheets'""",
        (total,),
    )
    cnx.commit()
    log.info("Planilla de soporte actualizada: %s filas", total)


def aplicar_archivo(ruta):
    hojas = parsear(ruta)
    cnx = conectar()
    try:
        asegurar_tablas(cnx)
        aplicar(cnx, hojas)
    except Exception:
        cnx.rollback()
        raise
    finally:
        cnx.close()


def sincronizar():
    destino = os.path.join(os.environ.get("TEMP", os.environ.get("TMP", "/tmp")), "stock-soporte.xlsx")
    descargar(sheet_id(), destino)
    aplicar_archivo(destino)
