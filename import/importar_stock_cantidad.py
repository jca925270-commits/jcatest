"""
importar_stock_cantidad.py
---------------------------
Importa el historial de STOCK_2026.xlsx (hojas ARGENTINA y PARAGUAY) a las
tablas `categorias_stock` y `movimientos_cantidad`. Es un import de UNA
SOLA VEZ (a diferencia de sync_watcher.py, que vigila cambios en vivo) —
correrlo de nuevo es seguro (no duplica), pero no está pensado para dejarlo
corriendo en segundo plano.

Uso:
    python importar_stock_cantidad.py "C:/ruta/a/STOCK_2026.xlsx"
"""

import configparser
import logging
import sys
import os
from datetime import datetime

import mysql.connector
from openpyxl import load_workbook

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("importar_stock_cantidad")

CONFIG = configparser.ConfigParser()
CONFIG.read(os.path.join(os.path.dirname(__file__), "config_sync.ini"))
DB_CFG = None
if CONFIG.has_section("db"):
    DB_CFG = {
        "host": CONFIG["db"]["host"],
        "database": CONFIG["db"]["database"],
        "user": CONFIG["db"]["user"],
        "password": CONFIG["db"]["password"],
    }

# Offset fijo observado en la planilla: la fila de encabezados de columna
# ("FECHA", "ESTADO", ... "EBD\nIngreso", ...) está 3 filas debajo de la
# fila que tiene los nombres de bloque/categoría (ej. "Oficina\n(EBD para Venta)").
OFFSET_FILA_BLOQUES = 3


def limpiar_texto(v):
    if v is None:
        return None
    return str(v).replace("\n", " ").strip() or None


def a_fecha(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    return None


def encontrar_fila_headers(ws):
    for num_fila, row in enumerate(ws.iter_rows(min_row=1, max_row=15), start=1):
        primera = row[0].value if len(row) else None
        if isinstance(primera, str) and primera.strip().upper() == "FECHA":
            return num_fila
    return None


def detectar_bloques(ws, fila_bloques, fila_headers):
    """Devuelve una lista de dicts: {nombre, tipo, col_ingreso, col_egreso}"""
    row_bloques = list(ws.iter_rows(min_row=fila_bloques, max_row=fila_bloques, values_only=True))[0]
    row_headers = list(ws.iter_rows(min_row=fila_headers, max_row=fila_headers, values_only=True))[0]

    bloques = []
    bloque_actual = None
    pendiente_ingreso = None  # (tipo, col) esperando su columna de egreso

    for col in range(8, len(row_headers) + 1):  # las primeras 7 columnas son metadata fija
        etiqueta_bloque = row_bloques[col - 1] if col - 1 < len(row_bloques) else None
        if etiqueta_bloque and str(etiqueta_bloque).strip().lower() != "o":
            bloque_actual = limpiar_texto(etiqueta_bloque)

        header = row_headers[col - 1] if col - 1 < len(row_headers) else None
        if not header or not isinstance(header, str):
            continue
        header_limpio = header.replace("\n", " ").strip()

        if "ingreso" in header_limpio.lower():
            tipo = header_limpio.split()[0]  # "EBD", "D3", "D3+", "Dongle", etc.
            pendiente_ingreso = {"nombre": bloque_actual, "tipo": tipo, "col_ingreso": col, "col_egreso": None}
        elif "egreso" in header_limpio.lower() and pendiente_ingreso is not None:
            pendiente_ingreso["col_egreso"] = col
            bloques.append(pendiente_ingreso)
            pendiente_ingreso = None

    return bloques


def obtener_categoria_id(cnx, cache, nombre, pais):
    clave = (nombre, pais)
    if clave not in cache:
        cur = cnx.cursor()
        cur.execute("SELECT id FROM categorias_stock WHERE nombre=%s AND pais=%s", (nombre, pais))
        row = cur.fetchone()
        if row:
            cache[clave] = row[0]
        else:
            cur.execute("INSERT INTO categorias_stock (nombre, pais) VALUES (%s,%s)", (nombre, pais))
            cnx.commit()
            cache[clave] = cur.lastrowid
    return cache[clave]


def procesar_hoja(cnx, ws, hoja, pais, cache_categorias):
    fila_headers = encontrar_fila_headers(ws)
    if not fila_headers:
        log.warning("No se encontró la fila de encabezados ('FECHA') en hoja %s, se omite", hoja)
        return 0, 0

    fila_bloques = fila_headers - OFFSET_FILA_BLOQUES
    bloques = detectar_bloques(ws, fila_bloques, fila_headers)
    log.info("Hoja '%s': %s categorías detectadas", hoja, len(bloques))

    ok = err = 0
    cur = cnx.cursor()

    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_headers + 1), start=fila_headers + 1):
        fecha = a_fecha(row[0].value if len(row) > 0 else None)
        estado = limpiar_texto(row[1].value if len(row) > 1 else None)
        remito = limpiar_texto(row[2].value if len(row) > 2 else None)
        factura = limpiar_texto(row[3].value if len(row) > 3 else None)
        razon_social = limpiar_texto(row[4].value if len(row) > 4 else None)
        ticket = limpiar_texto(row[5].value if len(row) > 5 else None)
        motivo = limpiar_texto(row[6].value if len(row) > 6 else None)

        for b in bloques:
            ci, ce = b["col_ingreso"] - 1, b["col_egreso"] - 1
            ingreso = row[ci].value if ci < len(row) else None
            egreso = row[ce].value if ce < len(row) else None
            if ingreso is None and egreso is None:
                continue
            try:
                cat_id = obtener_categoria_id(cnx, cache_categorias, b["nombre"] or b["tipo"], pais)
                cur.execute(
                    """INSERT IGNORE INTO movimientos_cantidad
                       (categoria_id, fecha, estado, remito, factura, razon_social, ticket, motivo,
                        ingreso, egreso, origen_hoja, origen_fila)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (cat_id, fecha, estado, remito, factura, razon_social, ticket, motivo,
                     ingreso, egreso, hoja, fila_num),
                )
                cnx.commit()
                ok += 1
            except mysql.connector.Error as e:
                cnx.rollback()
                log.error("Error fila %s, categoría '%s' (hoja %s): %s", fila_num, b["nombre"], hoja, e)
                err += 1

    return ok, err


def recalcular_totales(cnx):
    cur = cnx.cursor()
    cur.execute(
        """UPDATE categorias_stock c
           SET cantidad_actual = COALESCE((
               SELECT SUM(COALESCE(ingreso,0)) - SUM(COALESCE(egreso,0))
               FROM movimientos_cantidad m WHERE m.categoria_id = c.id
           ), 0)"""
    )
    cnx.commit()
    log.info("Totales recalculados para todas las categorías.")


def main():
    if len(sys.argv) < 2:
        print('Uso: python importar_stock_cantidad.py "C:/ruta/a/STOCK_2026.xlsx"')
        sys.exit(1)

    ruta = sys.argv[1]
    if not os.path.exists(ruta):
        log.error("No se encuentra el archivo: %s", ruta)
        sys.exit(1)

    if not DB_CFG:
        log.error("Falta import/config_sync.ini con la seccion [db]")
        sys.exit(1)

    wb = load_workbook(ruta, data_only=True, read_only=True)
    cnx = mysql.connector.connect(**DB_CFG)
    cache_categorias = {}

    mapa_paises = {"ARGENTINA": "ARGENTINA", "PARAGUAY": "PARAGUAY"}
    total_ok = total_err = 0

    for hoja in wb.sheetnames:
        pais = mapa_paises.get(hoja.upper(), hoja.upper())
        ok, err = procesar_hoja(cnx, wb[hoja], hoja, pais, cache_categorias)
        log.info("Hoja '%s': %s movimientos importados, %s errores", hoja, ok, err)
        total_ok += ok
        total_err += err

    recalcular_totales(cnx)

    cur = cnx.cursor()
    cur.execute("SELECT nombre, pais, cantidad_actual FROM categorias_stock ORDER BY pais, nombre")
    print("\n=== Stock actual por categoría ===")
    for nombre, pais, cantidad in cur.fetchall():
        print(f"  [{pais}] {nombre}: {cantidad}")

    cnx.close()
    log.info("Importación completa: %s movimientos ok, %s errores", total_ok, total_err)


if __name__ == "__main__":
    main()
