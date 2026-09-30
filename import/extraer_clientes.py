"""
extraer_clientes.py
---------------------
Extrae los nombres de cliente (razón social) del archivo
movimientos_soporte_tecnico_y_stock.xlsx y los carga como registros
reales en la tabla `clientes` — hasta ahora esos nombres solo vivían
como texto suelto dentro del detalle de cada movimiento.

Además, VINCULA retroactivamente los movimientos ya importados por
importar_soporte_tecnico.py con su cliente correspondiente (llenando
movimientos.cliente_id), para que se vea prolijo en Reportes y
Movimientos.

Fuentes:
  - Hoja "razon_social": lista de nombres ya cargados en el sistema
  - Hojas "Ingresos" y "Egresos": columna razon_social / razsocial_agreg
    de cada movimiento (esto es lo que permite el vínculo retroactivo)

Es un import de UNA SOLA VEZ. Correrlo de nuevo es seguro: no duplica
clientes (se dedupe por nombre exacto) y solo actualiza movimientos que
todavía no tuvieran cliente asignado.

Uso:
    python extraer_clientes.py "C:/ruta/al/movimientos soporte tecnico y stock.xlsx"
"""

import sys
import os
import logging

import mysql.connector
from openpyxl import load_workbook

from sync_watcher import Catalogos, DB_CFG
from importar_reparados import buscar_valor, encontrar_header

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("extraer_clientes")


def procesar_lista_razon_social(cat, ws):
    creados_o_existentes = 0
    vistos = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        valor = row[0] if row else None
        if not valor:
            continue
        nombre = str(valor).strip()
        if not nombre or nombre.lower() in ("razon social", "nombres"):
            continue
        if nombre in vistos:
            continue
        vistos.add(nombre)
        cat.cliente(nombre)
        creados_o_existentes += 1
    return creados_o_existentes


def procesar_hoja_movimientos(cnx, cat, ws, hoja):
    fila_header, cols = encontrar_header(ws)
    if not fila_header:
        log.warning("Hoja '%s': no se encontró columna 'Mac', se omite", hoja)
        return 0

    actualizados = 0
    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_header + 1), start=fila_header + 1):
        razon = buscar_valor(row, cols, "razon_social") or buscar_valor(row, cols, "razsocial_agreg")
        if not razon:
            continue
        nombre = str(razon).strip()
        if not nombre:
            continue

        cliente_id = cat.cliente(nombre)
        cur = cnx.cursor()
        cur.execute(
            """UPDATE movimientos SET cliente_id = %s
               WHERE origen_hoja = %s AND origen_fila = %s AND cliente_id IS NULL""",
            (cliente_id, hoja, fila_num),
        )
        actualizados += cur.rowcount
        cnx.commit()

    return actualizados


def main():
    if len(sys.argv) < 2:
        print('Uso: python extraer_clientes.py "C:/ruta/al/movimientos soporte tecnico y stock.xlsx"')
        sys.exit(1)
    ruta = sys.argv[1]
    if not os.path.exists(ruta):
        log.error("No se encuentra el archivo: %s", ruta)
        sys.exit(1)

    wb = load_workbook(ruta, data_only=True, read_only=True)
    cnx = mysql.connector.connect(**DB_CFG)
    cat = Catalogos(cnx)

    if "razon_social" in wb.sheetnames:
        n = procesar_lista_razon_social(cat, wb["razon_social"])
        log.info("Hoja 'razon_social': %s nombres procesados (creados o ya existentes)", n)

    total_actualizados = 0
    for hoja in ["Ingresos", "Egresos"]:
        if hoja not in wb.sheetnames:
            continue
        n = procesar_hoja_movimientos(cnx, cat, wb[hoja], hoja)
        log.info("Hoja '%s': %s movimientos vinculados a un cliente", hoja, n)
        total_actualizados += n

    cur = cnx.cursor()
    cur.execute("SELECT COUNT(*) FROM clientes")
    total_clientes = cur.fetchone()[0]

    cnx.close()
    log.info("Listo: %s clientes en total en la base, %s movimientos vinculados retroactivamente",
              total_clientes, total_actualizados)


if __name__ == "__main__":
    main()
