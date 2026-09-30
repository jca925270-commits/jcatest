"""
importar_soporte_tecnico.py
-----------------------------
Importa las hojas "Ingresos" y "Egresos" de
movimientos_soporte_tecnico_y_stock.xlsx hacia las tablas `dispositivos`
y `movimientos` — el MISMO esquema que usa toda la trazabilidad por MAC.

No se crean tablas nuevas: la MAC sigue siendo el eje. Cada ticket de
soporte queda como un movimiento más en el historial del dispositivo
correspondiente (visible también en su ficha individual), y además se
puede filtrar por separado en la sección "Soporte Técnico" del sistema
(gracias a que quedan marcados con origen_hoja = 'Ingresos'/'Egresos').

Es un import de UNA SOLA VEZ. Correrlo de nuevo es seguro (no duplica
movimientos), pero puede actualizar el estado de dispositivos existentes.

Uso:
    python importar_soporte_tecnico.py "C:/ruta/al/movimientos_soporte_tecnico_y_stock.xlsx"
"""

import sys
import os
import logging

import mysql.connector
from openpyxl import load_workbook

from sync_watcher import normalizar_mac, a_fecha, Catalogos, buscar_dispositivo, DB_CFG
from importar_reparados import extraer_macs, normalizar_tipo, buscar_valor, encontrar_header

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("importar_soporte_tecnico")


def determinar_estado(instancia, estado_texto, es_ingreso):
    inst = (instancia or "").strip().upper()

    if "NO INGRESO" in inst:
        return None  # no hubo movimiento físico real, se omite

    es_garantia = estado_texto and "garant" in str(estado_texto).lower()

    if "FINALIZADO" in inst:
        if es_garantia:
            return "garantia"
        return "reparado" if es_ingreso else "devuelto"

    # EN PROCESO, PENDIENTE, o cualquier otro valor no reconocido
    return "a_reparar"


def procesar_hoja(cnx, cat, ws, hoja, es_ingreso):
    fila_header, cols = encontrar_header(ws)
    if not fila_header:
        log.warning("Hoja '%s': no se encontró columna 'Mac', se omite", hoja)
        return 0, 0, 0

    ok = err = omitidas = 0
    tipo_mov_base = "a_reparar" if es_ingreso else "devolucion"

    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_header + 1), start=fila_header + 1):
        mac_texto = buscar_valor(row, cols, "Mac")
        mac_wifi, mac_eth = extraer_macs(mac_texto)
        if not mac_wifi and not mac_eth:
            continue  # fila sin MAC identificable (ej. solo un n° de seguimiento)

        razon_social = buscar_valor(row, cols, "razon_social") or buscar_valor(row, cols, "razsocial_agreg")
        tipo_texto = buscar_valor(row, cols, "disp_ingreso") or buscar_valor(row, cols, "disp_egreso")
        estado_texto = buscar_valor(row, cols, "Estado")
        instancia = buscar_valor(row, cols, "instancia_gestion")
        transportista = buscar_valor(row, cols, "correo_ingreso") or buscar_valor(row, cols, "Correo")
        seguimiento = buscar_valor(row, cols, "segui_ingreso") or buscar_valor(row, cols, "Seguimiento")
        ticket = buscar_valor(row, cols, "ticket")
        observacion = buscar_valor(row, cols, "observacion") or buscar_valor(row, cols, "obs_egreso")
        contacto = buscar_valor(row, cols, "con_ingreso") or buscar_valor(row, cols, "con_soporte_egreso")
        fecha = a_fecha(buscar_valor(row, cols, "Datetime") or buscar_valor(row, cols, "fecha_egreso"))

        estado_final = determinar_estado(instancia, estado_texto, es_ingreso)
        if estado_final is None:
            omitidas += 1
            continue

        tipo_id = cat.tipo_equipo(normalizar_tipo(tipo_texto))
        tipo_mov = "garantia" if estado_final == "garantia" else tipo_mov_base

        obs_final = " | ".join(filter(None, [
            f"Cliente: {razon_social}" if razon_social else None,
            f"Contacto: {contacto}" if contacto else None,
            str(observacion).strip() if observacion else None,
            f"Instancia: {instancia}" if instancia else None,
            f"[Soporte técnico - {hoja}]",
        ]))[:250]

        ticket_final = (str(ticket).strip() if ticket else None) or \
                       (str(seguimiento).strip()[:50] if seguimiento else None)

        try:
            cur = cnx.cursor()
            disp_id = buscar_dispositivo(cnx, mac_wifi, mac_eth)
            if disp_id:
                cur.execute(
                    """UPDATE dispositivos SET mac_wifi=COALESCE(%s, mac_wifi),
                       mac_eth=COALESCE(%s, mac_eth), estado=%s, observaciones=%s,
                       fecha_llegada=COALESCE(fecha_llegada, %s)
                       WHERE id=%s""",
                    (mac_wifi, mac_eth, estado_final, obs_final, fecha, disp_id),
                )
            else:
                cod_interno = "INT-" + (mac_wifi or mac_eth).replace(":", "")
                cur.execute(
                    """INSERT INTO dispositivos
                       (mac_wifi, mac_eth, tipo_equipo_id, estado, codigo_barras_interno,
                        observaciones, fecha_llegada)
                       VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                    (mac_wifi, mac_eth, tipo_id, estado_final, cod_interno, obs_final, fecha),
                )
                disp_id = cur.lastrowid
            cnx.commit()

            cur.execute(
                """INSERT IGNORE INTO movimientos
                   (dispositivo_id, tipo_movimiento, remito, ticket, motivo, origen_hoja, origen_fila)
                   VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                (disp_id, tipo_mov, (transportista or "")[:50] or None, ticket_final,
                 obs_final, hoja, fila_num),
            )
            cnx.commit()
            ok += 1
        except mysql.connector.Error as e:
            cnx.rollback()
            log.error("Error fila %s de %s: %s", fila_num, hoja, e)
            err += 1

    return ok, err, omitidas


def main():
    if len(sys.argv) < 2:
        print('Uso: python importar_soporte_tecnico.py "C:/ruta/al/movimientos_soporte_tecnico_y_stock.xlsx"')
        sys.exit(1)
    ruta = sys.argv[1]
    if not os.path.exists(ruta):
        log.error("No se encuentra el archivo: %s", ruta)
        sys.exit(1)

    wb = load_workbook(ruta, data_only=True, read_only=True)
    cnx = mysql.connector.connect(**DB_CFG)
    cat = Catalogos(cnx)

    total_ok = total_err = total_omitidas = 0
    for hoja, es_ingreso in [("Ingresos", True), ("Egresos", False)]:
        if hoja not in wb.sheetnames:
            log.warning("No se encontró la hoja '%s' en el archivo, se omite", hoja)
            continue
        ok, err, omitidas = procesar_hoja(cnx, cat, wb[hoja], hoja, es_ingreso)
        log.info("Hoja '%s': %s ok, %s errores, %s omitidas (NO INGRESO)", hoja, ok, err, omitidas)
        total_ok += ok
        total_err += err
        total_omitidas += omitidas

    cnx.close()
    log.info("Importación completa: %s ok, %s errores, %s omitidas", total_ok, total_err, total_omitidas)


if __name__ == "__main__":
    main()
