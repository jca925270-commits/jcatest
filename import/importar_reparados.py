"""
importar_reparados.py
-----------------------
Importa el historial de "Relevamiento - Equipos Reparados y Reutilizados
para VENTA.xlsx" hacia las tablas `dispositivos` y `movimientos` — el
MISMO esquema que usa la trazabilidad por MAC (a diferencia de
STOCK_2026.xlsx, este archivo sí tiene MAC por fila).

Cada hoja combina las dos MAC (WiFi + Ethernet) en una sola celda,
separadas por "|", espacios o tabulaciones — el script las separa
automáticamente.

Es un import de UNA SOLA VEZ. Correrlo de nuevo es seguro (no duplica
movimientos, gracias a origen_hoja+origen_fila), pero puede ACTUALIZAR
el estado de dispositivos ya existentes si los volvés a correr.

Uso:
    python importar_reparados.py "C:/ruta/al/Relevamiento...xlsx"
"""

import sys
import os
import re
import logging

import mysql.connector
from openpyxl import load_workbook

# Reutiliza utilidades ya probadas de sync_watcher.py (misma carpeta)
from sync_watcher import normalizar_mac, a_fecha, Catalogos, buscar_dispositivo, DB_CFG

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("importar_reparados")


def extraer_macs(texto):
    """Encuentra hasta 2 direcciones MAC en un texto, sin importar el separador usado."""
    if not texto:
        return None, None
    encontrados = re.findall(r"[0-9A-Fa-f]{2}(?:[:\-]?[0-9A-Fa-f]{2}){5}", str(texto))
    macs = [normalizar_mac(m) for m in encontrados]
    macs = [m for m in macs if m]
    mac_wifi = macs[0] if len(macs) >= 1 else None
    mac_eth = macs[1] if len(macs) >= 2 else None
    return mac_wifi, mac_eth


def normalizar_tipo(texto):
    if not texto:
        return "EBD"
    t = str(texto).lower()
    if "d3+" in t:
        return "D3+"
    if "d3" in t:
        return "D3"
    if "black" in t:
        return "Blackbox"
    if "cube" in t:
        return "CUBE"
    if "ebd" in t:
        return "EBD"
    return str(texto).strip()[:50]


def mapear_estado(texto, default="reparado"):
    if not texto:
        return default
    t = str(texto).lower()
    if "vendid" in t:
        return "vendido"
    if "comodat" in t:
        return "comodato"
    if "falla" in t:
        return "a_reparar"   # supuesto: "falla" = necesita reparación, no baja definitiva
    if "prueba" in t or "testeo" in t:
        return "prueba"
    if "finaliz" in t or "funcionamiento" in t or "listo" in t:
        return "reparado"
    if "devuel" in t:
        return "devuelto"
    return default


def buscar_valor(row, cols, *nombres):
    """Busca el valor de una celda por coincidencia parcial de nombre de columna
    (soporta encabezados con espacios extra, flechas, mayúsculas distintas, etc.)."""
    for n in nombres:
        n_norm = n.strip().lower()
        for clave, idx_col in cols.items():
            if n_norm in clave.strip().lower():
                idx = idx_col - 1
                if idx < len(row):
                    return row[idx].value
    return None


def encontrar_header(ws):
    """Busca la fila que tiene una columna llamada exactamente 'MAC'."""
    for num_fila, row in enumerate(ws.iter_rows(min_row=1, max_row=5), start=1):
        for c in row:
            valor = getattr(c, "value", None)
            if isinstance(valor, str) and valor.strip().lower() == "mac":
                mapa = {}
                for cc in row:
                    if isinstance(getattr(cc, "value", None), str) and cc.value.strip():
                        mapa[cc.value.strip()] = cc.column
                return num_fila, mapa
    return None, None


def procesar_hoja_generica(cnx, cat, ws, hoja, estado_default, tipo_movimiento_default, deposito_default=None):
    fila_header, cols = encontrar_header(ws)
    if not fila_header:
        log.warning("Hoja '%s': no se encontró columna 'MAC', se omite", hoja)
        return 0, 0

    ok = err = 0
    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_header + 1), start=fila_header + 1):
        mac_texto = buscar_valor(row, cols, "MAC")
        mac_wifi, mac_eth = extraer_macs(mac_texto)
        if not mac_wifi and not mac_eth:
            continue

        tipo_texto = buscar_valor(row, cols, "Tipo de Fuente", "FUENTE")
        tipo_id = cat.tipo_equipo(normalizar_tipo(tipo_texto))

        ubic_texto = buscar_valor(row, cols, "Depósito", "UBICACION", "ORIENTACION") or deposito_default or "Sin especificar"
        ubic_id = cat.deposito(str(ubic_texto).strip()[:100])

        estado_texto = buscar_valor(row, cols, "Estado", "ESTADO")
        estado_final = mapear_estado(estado_texto, estado_default)

        observaciones = buscar_valor(row, cols, "Observaciones")
        entidad = buscar_valor(row, cols, "Entidad", "ENTIDAD")
        fecha_inicio = a_fecha(buscar_valor(row, cols, "Fecha Comienzo Testeo"))

        obs_final = " | ".join(filter(None, [
            f"Entidad: {entidad}" if entidad else None,
            str(observaciones).strip() if observaciones else None,
            f"[Migrado desde '{hoja}']",
        ]))[:65000]

        try:
            cur = cnx.cursor()
            disp_id = buscar_dispositivo(cnx, mac_wifi, mac_eth)
            if disp_id:
                cur.execute(
                    """UPDATE dispositivos SET mac_wifi=COALESCE(%s, mac_wifi),
                       mac_eth=COALESCE(%s, mac_eth), deposito_actual_id=%s,
                       estado=%s, observaciones=%s, fecha_llegada=COALESCE(fecha_llegada, %s)
                       WHERE id=%s""",
                    (mac_wifi, mac_eth, ubic_id, estado_final, obs_final, fecha_inicio, disp_id),
                )
            else:
                cod_interno = "INT-" + (mac_wifi or mac_eth).replace(":", "")
                cur.execute(
                    """INSERT INTO dispositivos
                       (mac_wifi, mac_eth, tipo_equipo_id, deposito_actual_id, estado,
                        codigo_barras_interno, observaciones, fecha_llegada)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (mac_wifi, mac_eth, tipo_id, ubic_id, estado_final, cod_interno, obs_final, fecha_inicio),
                )
                disp_id = cur.lastrowid
            cnx.commit()

            cur.execute(
                """INSERT IGNORE INTO movimientos
                   (dispositivo_id, tipo_movimiento, motivo, origen_hoja, origen_fila)
                   VALUES (%s,%s,%s,%s,%s)""",
                (disp_id, tipo_movimiento_default, f"Migrado desde '{hoja}' (Relevamiento reparados)", hoja, fila_num),
            )
            cnx.commit()
            ok += 1
        except mysql.connector.Error as e:
            cnx.rollback()
            log.error("Error fila %s de %s: %s", fila_num, hoja, e)
            err += 1

    return ok, err


def procesar_cube_showroom(cnx, cat, ws, hoja):
    """Hoja simple: solo MAC + cantidad de pantallas, sin columnas de estado."""
    ok = err = 0
    tipo_id = cat.tipo_equipo("CUBE")
    ubic_id = cat.deposito("Showroom")
    for fila_num, row in enumerate(ws.iter_rows(min_row=2), start=2):
        mac_texto = row[0].value if len(row) else None
        mac_wifi, mac_eth = extraer_macs(mac_texto)
        if not mac_wifi and not mac_eth:
            continue
        try:
            cur = cnx.cursor()
            disp_id = buscar_dispositivo(cnx, mac_wifi, mac_eth)
            if disp_id:
                cur.execute("UPDATE dispositivos SET deposito_actual_id=%s, estado='reparado' WHERE id=%s",
                            (ubic_id, disp_id))
            else:
                cod_interno = "INT-" + (mac_wifi or mac_eth).replace(":", "")
                cur.execute(
                    """INSERT INTO dispositivos
                       (mac_wifi, mac_eth, tipo_equipo_id, deposito_actual_id, estado, codigo_barras_interno)
                       VALUES (%s,%s,%s,%s,'reparado',%s)""",
                    (mac_wifi, mac_eth, tipo_id, ubic_id, cod_interno),
                )
                disp_id = cur.lastrowid
            cnx.commit()
            cur.execute(
                """INSERT IGNORE INTO movimientos
                   (dispositivo_id, tipo_movimiento, motivo, origen_hoja, origen_fila)
                   VALUES (%s,'reparado','Migrado desde CUBE REPARADOS SHOWROOM',%s,%s)""",
                (disp_id, hoja, fila_num),
            )
            cnx.commit()
            ok += 1
        except mysql.connector.Error as e:
            cnx.rollback()
            log.error("Error fila %s de %s: %s", fila_num, hoja, e)
            err += 1
    return ok, err


# Por hoja: (estado por defecto si la fila no trae uno claro, tipo de movimiento a registrar, depósito por defecto)
HOJAS_CONFIG = {
    "Testeo EBD-Reparados": ("reparado", "reparado", None),
    "D3refresh": ("reparado", "reparado", None),
    "No pasaron prueba": ("a_reparar", "a_reparar", None),
    "para PERU": ("reparado", "reparado", "Peru"),
    "comodato EBD": ("comodato", "comodato", None),
    "CUBE Devueltos": ("en_stock", "devolucion", None),
    "Blackbox Devueltos": ("en_stock", "devolucion", None),
    "D3 Falla": ("a_reparar", "a_reparar", None),
}


def main():
    if len(sys.argv) < 2:
        print('Uso: python importar_reparados.py "C:/ruta/al/Relevamiento...xlsx"')
        sys.exit(1)
    ruta = sys.argv[1]
    if not os.path.exists(ruta):
        log.error("No se encuentra el archivo: %s", ruta)
        sys.exit(1)

    wb = load_workbook(ruta, data_only=True, read_only=True)
    cnx = mysql.connector.connect(**DB_CFG)
    cat = Catalogos(cnx)

    total_ok = total_err = 0
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        if hoja == "Hoja 10":
            log.info("Hoja '%s' está vacía, se omite", hoja)
            continue
        if hoja == "CUBE REPARADOS SHOWROOM":
            ok, err = procesar_cube_showroom(cnx, cat, ws, hoja)
        elif hoja in HOJAS_CONFIG:
            estado_default, tipo_mov, ubic_default = HOJAS_CONFIG[hoja]
            ok, err = procesar_hoja_generica(cnx, cat, ws, hoja, estado_default, tipo_mov, ubic_default)
        else:
            log.info("Hoja '%s' no reconocida, se omite", hoja)
            continue
        log.info("Hoja '%s': %s ok, %s errores", hoja, ok, err)
        total_ok += ok
        total_err += err

    cnx.close()
    log.info("Importación completa: %s ok, %s errores", total_ok, total_err)


if __name__ == "__main__":
    main()
