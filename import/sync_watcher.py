"""
sync_watcher.py
----------------
Vigila BASE_MAC_ADDRESS-2025-2026.xlsx y, cada vez que se guarda un cambio,
sincroniza automáticamente los dispositivos y movimientos hacia MySQL.

Uso:
    pip install -r requirements.txt
    python sync_watcher.py

Para dejarlo corriendo siempre en el servidor, ver la sección de systemd
en el README (import/README.md).
"""

import configparser
import logging
import os
import re
import threading
import time
from datetime import datetime

import mysql.connector
from openpyxl import load_workbook
from watchdog.events import FileSystemEventHandler
from watchdog.observers import Observer

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
log = logging.getLogger("sync_watcher")

CONFIG = configparser.ConfigParser()
CONFIG.read(os.path.join(os.path.dirname(__file__), "config_sync.ini"))

DB_CFG = {
    "host": CONFIG["db"]["host"],
    "database": CONFIG["db"]["database"],
    "user": CONFIG["db"]["user"],
    "password": CONFIG["db"]["password"],
}
CARPETA = CONFIG["watcher"]["carpeta"]
ARCHIVO_MAC = CONFIG["watcher"]["archivo_mac"]
DEBOUNCE = float(CONFIG["watcher"].get("debounce_segundos", 4))


# ------------------------------------------------------------------
# Utilidades
# ------------------------------------------------------------------
def normalizar_mac(valor):
    if valor is None:
        return None
    limpio = re.sub(r"[^0-9A-Fa-f]", "", str(valor)).upper()
    if not limpio:
        return None
    if len(limpio) != 12:
        return limpio  # se guarda tal cual, no matchea formato estándar
    return ":".join(limpio[i:i + 2] for i in range(0, 12, 2))


def a_fecha(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    return None  # texto no parseable (ej. "Retiró 22/05/2025") -> se ignora la fecha


def encontrar_header(ws, columnas_clave):
    """Busca la fila que contiene todas las columnas_clave y devuelve (fila_idx, {nombre: col_idx})."""
    for num_fila, row in enumerate(ws.iter_rows(min_row=1, max_row=10), start=1):
        valores = {(c.value or "").strip() if isinstance(c.value, str) else c.value: c.column
                   for c in row if getattr(c, "value", None)}
        if all(any(k.lower() in str(v).lower() for v in valores if v) for k in columnas_clave):
            mapa = {}
            for c in row:
                if isinstance(getattr(c, "value", None), str) and c.value.strip():
                    mapa[c.value.strip()] = c.column
            return num_fila, mapa
    return None, None


# ------------------------------------------------------------------
# Acceso a catálogos (con caché en memoria durante una corrida)
# ------------------------------------------------------------------
class Catalogos:
    def __init__(self, cnx):
        self.cnx = cnx
        self.tipos = {}
        self.depositos = {}
        self.lotes = {}
        self.clientes = {}

    def tipo_equipo(self, nombre):
        nombre = nombre.strip()
        if nombre not in self.tipos:
            cur = self.cnx.cursor()
            cur.execute("SELECT id FROM tipos_equipo WHERE nombre = %s", (nombre,))
            row = cur.fetchone()
            if not row:
                cur.execute("INSERT INTO tipos_equipo (nombre) VALUES (%s)", (nombre,))
                self.cnx.commit()
                self.tipos[nombre] = cur.lastrowid
            else:
                self.tipos[nombre] = row[0]
        return self.tipos[nombre]

    def deposito(self, nombre):
        nombre = nombre.strip()
        if nombre not in self.depositos:
            cur = self.cnx.cursor()
            cur.execute("SELECT id FROM depositos WHERE nombre = %s", (nombre,))
            row = cur.fetchone()
            if not row:
                cur.execute("INSERT INTO depositos (nombre) VALUES (%s)", (nombre,))
                self.cnx.commit()
                self.depositos[nombre] = cur.lastrowid
            else:
                self.depositos[nombre] = row[0]
        return self.depositos[nombre]

    def lote(self, nombre):
        if not nombre:
            return None
        nombre = str(nombre).strip()
        if nombre not in self.lotes:
            cur = self.cnx.cursor()
            cur.execute("SELECT id FROM lotes WHERE nombre = %s", (nombre,))
            row = cur.fetchone()
            if not row:
                cur.execute("INSERT INTO lotes (nombre) VALUES (%s)", (nombre,))
                self.cnx.commit()
                self.lotes[nombre] = cur.lastrowid
            else:
                self.lotes[nombre] = row[0]
        return self.lotes[nombre]

    def cliente(self, razon_social, email=None, entidad=None, sucursal=None):
        if not razon_social:
            return None
        razon_social = str(razon_social).strip()
        if razon_social not in self.clientes:
            cur = self.cnx.cursor()
            cur.execute("SELECT id FROM clientes WHERE razon_social = %s", (razon_social,))
            row = cur.fetchone()
            if not row:
                cur.execute(
                    "INSERT INTO clientes (razon_social, email, entidad, sucursal) VALUES (%s,%s,%s,%s)",
                    (razon_social, email, entidad, sucursal),
                )
                self.cnx.commit()
                self.clientes[razon_social] = cur.lastrowid
            else:
                self.clientes[razon_social] = row[0]
        return self.clientes[razon_social]


def buscar_dispositivo(cnx, mac_wifi, mac_eth):
    cur = cnx.cursor()
    if mac_wifi:
        cur.execute("SELECT id FROM dispositivos WHERE mac_wifi = %s", (mac_wifi,))
        row = cur.fetchone()
        if row:
            return row[0]
    if mac_eth:
        cur.execute("SELECT id FROM dispositivos WHERE mac_eth = %s", (mac_eth,))
        row = cur.fetchone()
        if row:
            return row[0]
    return None


# ------------------------------------------------------------------
# Procesadores por tipo de hoja
# ------------------------------------------------------------------
def procesar_ingreso(cnx, cat, ws, hoja, tipo_nombre, deposito_nombre):
    fila_header, cols = encontrar_header(ws, ["MAC WIFI"])
    if not fila_header:
        log.warning("No se encontró encabezado en hoja %s, se omite", hoja)
        return 0, 0

    ok = err = 0
    tipo_id = cat.tipo_equipo(tipo_nombre)
    ubic_id = cat.deposito(deposito_nombre)

    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_header + 1), start=fila_header + 1):
        get = lambda nombre: (row[cols[nombre] - 1].value if (nombre in cols and cols[nombre] - 1 < len(row)) else None)

        mac_wifi = normalizar_mac(get("MAC WIFI"))
        mac_eth = normalizar_mac(get("MAC ETH"))
        if not mac_wifi and not mac_eth:
            continue  # fila vacía

        lote_id = cat.lote(get("Lote"))
        caja = str(get("Caja n°") or get("Caja n") or "").strip() or None
        fecha_llegada = a_fecha(get("Fecha de llegada"))
        fecha_carga = a_fecha(get("Fecha de carga"))

        try:
            cur = cnx.cursor()
            disp_id = buscar_dispositivo(cnx, mac_wifi, mac_eth)
            if disp_id:
                cur.execute(
                    """UPDATE dispositivos SET mac_wifi=COALESCE(%s, mac_wifi),
                       mac_eth=COALESCE(%s, mac_eth), lote_id=%s, caja_n=%s,
                       deposito_actual_id=%s, fecha_llegada=%s, fecha_carga=%s
                       WHERE id=%s""",
                    (mac_wifi, mac_eth, lote_id, caja, ubic_id, fecha_llegada, fecha_carga, disp_id),
                )
            else:
                cur.execute(
                    """INSERT INTO dispositivos
                       (mac_wifi, mac_eth, tipo_equipo_id, lote_id, caja_n,
                        deposito_actual_id, estado, codigo_barras_interno, fecha_llegada, fecha_carga)
                       VALUES (%s,%s,%s,%s,%s,%s,'en_stock',%s,%s,%s)""",
                    (mac_wifi, mac_eth, tipo_id, lote_id, caja, ubic_id,
                     "INT-" + (mac_wifi or mac_eth).replace(":", ""), fecha_llegada, fecha_carga),
                )
                disp_id = cur.lastrowid
            cnx.commit()

            # Movimiento de ingreso (idempotente por hoja+fila)
            cur.execute(
                """INSERT IGNORE INTO movimientos
                   (dispositivo_id, tipo_movimiento, deposito_destino_id, motivo, origen_hoja, origen_fila)
                   VALUES (%s,'ingreso',%s,'Sincronizado desde Excel',%s,%s)""",
                (disp_id, ubic_id, hoja, fila_num),
            )
            cnx.commit()
            ok += 1
        except mysql.connector.Error as e:
            cnx.rollback()
            log.error("Error fila %s de %s: %s", fila_num, hoja, e)
            err += 1

    return ok, err


def procesar_egreso(cnx, cat, ws, hoja, movimiento_default="venta"):
    fila_header, cols = encontrar_header(ws, ["MAC WIFI", "Cliente"])
    if not fila_header:
        log.warning("No se encontró encabezado en hoja %s, se omite", hoja)
        return 0, 0

    ok = err = 0
    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_header + 1), start=fila_header + 1):
        get = lambda nombre: (row[cols[nombre] - 1].value if (nombre in cols and cols[nombre] - 1 < len(row)) else None)

        mac_wifi = normalizar_mac(get("MAC WIFI"))
        mac_eth = normalizar_mac(get("MAC ETH"))
        cliente_txt = get("Cliente")
        if not (mac_wifi or mac_eth) or not cliente_txt:
            continue

        try:
            disp_id = buscar_dispositivo(cnx, mac_wifi, mac_eth)
            if not disp_id:
                log.warning("Egreso de MAC no encontrada en stock (hoja %s, fila %s): %s/%s",
                            hoja, fila_num, mac_wifi, mac_eth)
                err += 1
                continue

            cliente_id = cat.cliente(cliente_txt)
            tipo_mov = "comodato" if "comodato" in hoja.lower() else movimiento_default

            cur = cnx.cursor()
            cur.execute(
                """INSERT IGNORE INTO movimientos
                   (dispositivo_id, tipo_movimiento, cliente_id, motivo, origen_hoja, origen_fila)
                   VALUES (%s,%s,%s,'Sincronizado desde Excel',%s,%s)""",
                (disp_id, tipo_mov, cliente_id, hoja, fila_num),
            )
            if cur.rowcount:  # solo actualiza el estado si es un movimiento nuevo
                estado = "comodato" if tipo_mov == "comodato" else "vendido"
                cur.execute(
                    "UPDATE dispositivos SET estado=%s, cliente_actual_id=%s WHERE id=%s",
                    (estado, cliente_id, disp_id),
                )
            cnx.commit()
            ok += 1
        except mysql.connector.Error as e:
            cnx.rollback()
            log.error("Error fila %s de %s: %s", fila_num, hoja, e)
            err += 1

    return ok, err


def procesar_devolucion(cnx, cat, ws, hoja):
    fila_header, cols = encontrar_header(ws, ["MAC WIFI", "Devolucion"])
    if not fila_header:
        return 0, 0

    ok = err = 0
    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_header + 1), start=fila_header + 1):
        get = lambda nombre: (row[cols[nombre] - 1].value if (nombre in cols and cols[nombre] - 1 < len(row)) else None)
        mac_wifi = normalizar_mac(get("MAC WIFI"))
        mac_eth = normalizar_mac(get("MAC ETH"))
        if not (mac_wifi or mac_eth):
            continue
        try:
            disp_id = buscar_dispositivo(cnx, mac_wifi, mac_eth)
            if not disp_id:
                err += 1
                continue
            cur = cnx.cursor()
            cur.execute(
                """INSERT IGNORE INTO movimientos
                   (dispositivo_id, tipo_movimiento, motivo, origen_hoja, origen_fila)
                   VALUES (%s,'devolucion','Sincronizado desde Excel',%s,%s)""",
                (disp_id, hoja, fila_num),
            )
            if cur.rowcount:
                cur.execute(
                    "UPDATE dispositivos SET estado='en_stock', cliente_actual_id=NULL WHERE id=%s",
                    (disp_id,),
                )
            cnx.commit()
            ok += 1
        except mysql.connector.Error as e:
            cnx.rollback()
            log.error("Error fila %s de %s: %s", fila_num, hoja, e)
            err += 1
    return ok, err


def procesar_migrados(cnx, cat, ws, hoja):
    fila_header, cols = encontrar_header(ws, ["MAC WIFI", "Razón Social"])
    if not fila_header:
        return 0, 0

    ok = err = 0
    for fila_num, row in enumerate(ws.iter_rows(min_row=fila_header + 1), start=fila_header + 1):
        get = lambda nombre: (row[cols[nombre] - 1].value if (nombre in cols and cols[nombre] - 1 < len(row)) else None)
        mac_wifi = normalizar_mac(get("MAC WIFI"))
        mac_eth = normalizar_mac(get("MAC ETH"))
        if not (mac_wifi or mac_eth):
            continue
        try:
            cliente_id = cat.cliente(get("Razón Social"), get("Email"), get("Entidad"), get("Sucursal"))
            disp_id = buscar_dispositivo(cnx, mac_wifi, mac_eth)
            cur = cnx.cursor()
            if not disp_id:
                tipo_id = cat.tipo_equipo(str(get("Tipo de Equipo") or "D3"))
                cur.execute(
                    """INSERT INTO dispositivos
                       (mac_wifi, mac_eth, tipo_equipo_id, estado, cliente_actual_id, codigo_barras_interno)
                       VALUES (%s,%s,%s,'vendido',%s,%s)""",
                    (mac_wifi, mac_eth, tipo_id, cliente_id, "INT-" + (mac_wifi or mac_eth).replace(":", "")),
                )
                disp_id = cur.lastrowid
                cnx.commit()
            cur.execute(
                """INSERT IGNORE INTO movimientos
                   (dispositivo_id, tipo_movimiento, cliente_id, motivo, origen_hoja, origen_fila)
                   VALUES (%s,'venta',%s,'Migración a 2.0 (sync Excel)',%s,%s)""",
                (disp_id, cliente_id, hoja, fila_num),
            )
            if cur.rowcount:
                cur.execute(
                    "UPDATE dispositivos SET estado='vendido', cliente_actual_id=%s WHERE id=%s",
                    (cliente_id, disp_id),
                )
            cnx.commit()
            ok += 1
        except mysql.connector.Error as e:
            cnx.rollback()
            log.error("Error fila %s de %s: %s", fila_num, hoja, e)
            err += 1
    return ok, err


def tipo_desde_nombre_hoja(hoja):
    h = hoja.lower()
    if "d3+" in h:
        return "D3+"
    if "d3" in h:
        return "D3"
    return "EBD"


def deposito_desde_nombre_hoja(hoja):
    h = hoja.lower()
    if "contenedor" in h:
        return "Contenedor"
    if "tienda nube" in h:
        return "Tienda Nube"
    return "Oficina"


# ------------------------------------------------------------------
# Sincronización completa de un archivo
# ------------------------------------------------------------------
def sincronizar(path_excel):
    log.info("Sincronizando %s ...", path_excel)
    wb = load_workbook(path_excel, data_only=True, read_only=True)
    cnx = mysql.connector.connect(**DB_CFG)
    cat = Catalogos(cnx)

    total_ok = total_err = 0
    detalle = []

    for hoja in wb.sheetnames:
        ws = wb[hoja]
        h = hoja.lower()
        if h.startswith("ingreso"):
            ok, err = procesar_ingreso(cnx, cat, ws, hoja,
                                        tipo_desde_nombre_hoja(hoja),
                                        deposito_desde_nombre_hoja(hoja))
        elif h.startswith("egreso"):
            ok, err = procesar_egreso(cnx, cat, ws, hoja)
        elif h.startswith("devo"):
            ok, err = procesar_devolucion(cnx, cat, ws, hoja)
        elif "migrad" in h:
            ok, err = procesar_migrados(cnx, cat, ws, hoja)
        else:
            log.info("Hoja '%s' no reconocida, se omite", hoja)
            continue

        total_ok += ok
        total_err += err
        detalle.append(f"{hoja}: {ok} ok, {err} errores")
        log.info("Hoja '%s': %s filas ok, %s errores", hoja, ok, err)

    cur = cnx.cursor()
    cur.execute(
        """INSERT INTO sync_control (archivo, ultima_sync, filas_ok, filas_error, detalle)
           VALUES (%s, NOW(), %s, %s, %s)
           ON DUPLICATE KEY UPDATE ultima_sync=NOW(), filas_ok=%s, filas_error=%s, detalle=%s""",
        (os.path.basename(path_excel), total_ok, total_err, "\n".join(detalle),
         total_ok, total_err, "\n".join(detalle)),
    )
    cnx.commit()
    cnx.close()
    log.info("Sincronización completa: %s ok, %s errores", total_ok, total_err)


# ------------------------------------------------------------------
# Watcher con debounce
# ------------------------------------------------------------------
class ExcelHandler(FileSystemEventHandler):
    def __init__(self, archivo_objetivo, callback, debounce):
        self.archivo_objetivo = archivo_objetivo
        self.callback = callback
        self.debounce = debounce
        self._timer = None
        self._lock = threading.Lock()

    def on_modified(self, event):
        if event.is_directory:
            return
        if os.path.basename(event.src_path) != self.archivo_objetivo:
            return
        with self._lock:
            if self._timer:
                self._timer.cancel()
            self._timer = threading.Timer(self.debounce, self._ejecutar, args=[event.src_path])
            self._timer.start()

    def _ejecutar(self, path):
        try:
            self.callback(path)
        except Exception:
            log.exception("Fallo al sincronizar %s", path)


def main():
    path_completo = os.path.join(CARPETA, ARCHIVO_MAC)
    if not os.path.exists(path_completo):
        log.error("No se encuentra el archivo: %s", path_completo)
        return

    # Sincronización inicial al arrancar
    sincronizar(path_completo)

    handler = ExcelHandler(ARCHIVO_MAC, sincronizar, DEBOUNCE)
    observer = Observer()
    observer.schedule(handler, CARPETA, recursive=False)
    observer.start()
    log.info("Vigilando cambios en %s ...", path_completo)

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
    observer.join()


if __name__ == "__main__":
    main()
