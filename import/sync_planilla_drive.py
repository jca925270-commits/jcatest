"""
Sincroniza la planilla STOCK 2026 desde Google Sheets hacia MySQL.

La hoja se edita en Drive. Este proceso la descarga, y si cambió, actualiza
categorias_stock, movimientos_cantidad y las razones sociales de la planilla.
El stock visible es la fila de números que está justo debajo de los nombres
de categoría (no la suma de movimientos).

Uso:
    python sync_planilla_drive.py
    python sync_planilla_drive.py --probar "C:/ruta/STOCK 2026.xlsx"
"""

import hashlib
import logging
import os
import re
import sys
import tempfile
import time
import urllib.error
import urllib.request

import mysql.connector
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
from importar_stock_cantidad import (  # noqa: E402
    a_fecha,
    detectar_bloques,
    encontrar_fila_headers,
    limpiar_texto,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("sync_planilla_drive")

HOJAS = {"ARGENTINA", "PARAGUAY"}
SUFIJO_EQUIPO = re.compile(r"\(\s*(?:EBD|D3\+?)\s+con\s+2[.,]0(?:\s+Lite)?\s*\)", re.I)
AVISO_COMPARTIR = (
    "Google no deja leer la planilla (sigue privada). "
    "En la hoja: Compartir, Acceso general, Cualquier persona con el enlace, Lector. "
    "Reintento en unos segundos."
)


CLAVES_ENTORNO = (
    "DB_HOST", "DB_PORT", "DB_NAME", "DB_USER", "DB_PASSWORD",
    "DB_SSL", "DB_SSL_CA", "DB_SSL_CA_PATH",
    "GOOGLE_SHEET_ID", "SYNC_INTERVAL_SEC",
)


def leer_env():
    ruta = os.path.join(os.path.dirname(__file__), "..", ".env")
    env = {}
    if os.path.exists(ruta):
        with open(ruta, encoding="utf-8") as archivo:
            for linea in archivo:
                linea = linea.strip()
                if not linea or linea.startswith("#") or "=" not in linea:
                    continue
                clave, valor = linea.split("=", 1)
                env[clave.strip()] = valor.strip()
    for clave in CLAVES_ENTORNO:
        if os.environ.get(clave):
            env[clave] = os.environ[clave]
    return env


def a_entero(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return int(valor)
    try:
        return int(float(str(valor).strip().replace(",", ".")))
    except ValueError:
        return None


def limpiar_razon(valor):
    texto = limpiar_texto(valor)
    if not texto:
        return None
    texto = SUFIJO_EQUIPO.sub("", texto)
    texto = re.sub(r"[ \t]+", " ", texto).strip()
    return texto or None


def recortar(valor, largo):
    if valor is None:
        return None
    return valor[:largo]


def parsear_libro(ruta):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hojas = []
    try:
        for nombre_hoja in libro.sheetnames:
            pais = nombre_hoja.strip().upper()
            if pais not in HOJAS:
                continue
            hoja = libro[nombre_hoja]
            fila_headers = encontrar_fila_headers(hoja)
            if not fila_headers:
                log.warning("La hoja %s no tiene la fila FECHA", nombre_hoja)
                continue
            bloques = detectar_bloques(hoja, fila_headers - 3, fila_headers)
            if pais == "ARGENTINA" and len(bloques) < 10:
                raise RuntimeError(f"La hoja ARGENTINA llegó incompleta ({len(bloques)} categorías)")
            if not bloques:
                raise RuntimeError(f"La hoja {pais} no tiene categorías")

            fila_stock = list(hoja.iter_rows(min_row=fila_headers - 2, max_row=fila_headers - 2, values_only=True))[0]
            categorias = []
            for bloque in bloques:
                indice = bloque["col_ingreso"] - 1
                stock = a_entero(fila_stock[indice] if indice < len(fila_stock) else None)
                categorias.append({"nombre": bloque["nombre"] or bloque["tipo"], "stock": stock})

            movimientos = []
            for fila_num, row in enumerate(hoja.iter_rows(min_row=fila_headers + 1), start=fila_headers + 1):
                fecha = a_fecha(row[0].value if len(row) > 0 else None)
                estado = recortar(limpiar_texto(row[1].value if len(row) > 1 else None), 50)
                remito = recortar(limpiar_texto(row[2].value if len(row) > 2 else None), 50)
                factura = recortar(limpiar_texto(row[3].value if len(row) > 3 else None), 50)
                razon = recortar(limpiar_razon(row[4].value if len(row) > 4 else None), 255)
                ticket = recortar(limpiar_texto(row[5].value if len(row) > 5 else None), 50)
                motivo = recortar(limpiar_texto(row[6].value if len(row) > 6 else None), 255)
                for bloque in bloques:
                    ingreso = a_entero(row[bloque["col_ingreso"] - 1].value if bloque["col_ingreso"] - 1 < len(row) else None)
                    egreso = a_entero(row[bloque["col_egreso"] - 1].value if bloque["col_egreso"] - 1 < len(row) else None)
                    if ingreso is None and egreso is None:
                        continue
                    movimientos.append({
                        "nombre": bloque["nombre"] or bloque["tipo"],
                        "fecha": fecha,
                        "estado": estado,
                        "remito": remito,
                        "factura": factura,
                        "razon": razon,
                        "ticket": ticket,
                        "motivo": motivo,
                        "ingreso": ingreso,
                        "egreso": egreso,
                        "fila": fila_num,
                    })
            if pais == "ARGENTINA" and not movimientos:
                raise RuntimeError("La hoja ARGENTINA no trajo movimientos")
            hojas.append({
                "hoja": nombre_hoja,
                "pais": pais,
                "categorias": categorias,
                "movimientos": movimientos,
            })
    finally:
        libro.close()
    if not hojas:
        raise RuntimeError("El archivo no tiene las hojas ARGENTINA y PARAGUAY")
    return hojas


def categoria_id(cur, cache, nombre, pais):
    clave = (nombre, pais)
    if clave not in cache:
        cur.execute("SELECT id FROM categorias_stock WHERE nombre=%s AND pais=%s", clave)
        fila = cur.fetchone()
        if fila:
            cache[clave] = fila[0]
        else:
            cur.execute("INSERT INTO categorias_stock (nombre, pais) VALUES (%s,%s)", clave)
            cache[clave] = cur.lastrowid
    return cache[clave]


def aplicar(cnx, hojas):
    cur = cnx.cursor()
    cache = {}
    razones = set()
    total = 0
    for hoja in hojas:
        vistos = []
        for mov in hoja["movimientos"]:
            cat = categoria_id(cur, cache, mov["nombre"], hoja["pais"])
            cur.execute(
                """INSERT INTO movimientos_cantidad
                   (categoria_id, fecha, estado, remito, factura, razon_social, ticket, motivo,
                    ingreso, egreso, origen_hoja, origen_fila)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                   ON DUPLICATE KEY UPDATE
                     fecha=VALUES(fecha), estado=VALUES(estado), remito=VALUES(remito),
                     factura=VALUES(factura), razon_social=VALUES(razon_social),
                     ticket=VALUES(ticket), motivo=VALUES(motivo),
                     ingreso=VALUES(ingreso), egreso=VALUES(egreso)""",
                (cat, mov["fecha"], mov["estado"], mov["remito"], mov["factura"], mov["razon"],
                 mov["ticket"], mov["motivo"], mov["ingreso"], mov["egreso"], hoja["hoja"], mov["fila"]),
            )
            vistos.append((cat, mov["fila"]))
            if mov["razon"]:
                razones.add(mov["razon"][:200])
            total += 1

        cur.execute(
            """CREATE TEMPORARY TABLE IF NOT EXISTS tmp_planilla_sync (
                 categoria_id INT NOT NULL,
                 origen_fila INT NOT NULL,
                 PRIMARY KEY (categoria_id, origen_fila)
               )"""
        )
        cur.execute("DELETE FROM tmp_planilla_sync")
        cur.executemany(
            "INSERT IGNORE INTO tmp_planilla_sync (categoria_id, origen_fila) VALUES (%s,%s)",
            vistos,
        )
        cur.execute(
            """DELETE m FROM movimientos_cantidad m
               LEFT JOIN tmp_planilla_sync t
                 ON t.categoria_id = m.categoria_id AND t.origen_fila = m.origen_fila
               WHERE m.origen_hoja = %s AND t.categoria_id IS NULL""",
            (hoja["hoja"],),
        )

        for categoria in hoja["categorias"]:
            if categoria["stock"] is None:
                continue
            cat = categoria_id(cur, cache, categoria["nombre"], hoja["pais"])
            cur.execute(
                "UPDATE categorias_stock SET cantidad_actual=%s WHERE id=%s",
                (categoria["stock"], cat),
            )

    for razon in razones:
        cur.execute("SELECT id FROM clientes WHERE razon_social=%s LIMIT 1", (razon,))
        if cur.fetchone():
            cur.execute("UPDATE clientes SET planilla=1 WHERE razon_social=%s AND planilla=0", (razon,))
        else:
            cur.execute("INSERT INTO clientes (razon_social, planilla) VALUES (%s,1)", (razon,))
    cur.execute(
        """INSERT INTO sync_control (archivo, ultima_sync, filas_ok, filas_error, detalle)
           VALUES ('STOCK_2026', NOW(), %s, 0, 'Google Sheets')
           ON DUPLICATE KEY UPDATE
             ultima_sync = NOW(),
             filas_ok = VALUES(filas_ok),
             filas_error = 0,
             detalle = 'Google Sheets'""",
        (total,),
    )
    cnx.commit()
    log.info("Planilla actualizada: %s movimientos en %s hojas", total, len(hojas))


def conectar():
    env = leer_env()
    host = env.get("DB_HOST", "127.0.0.1")
    config = {
        "host": host,
        "port": int(env.get("DB_PORT", "3306")),
        "database": env.get("DB_NAME", "stock_db"),
        "user": env.get("DB_USER", "root"),
        "password": env.get("DB_PASSWORD", ""),
    }
    if env.get("DB_SSL", "").lower() == "true" or host.endswith(".aivencloud.com"):
        ca_inline = env.get("DB_SSL_CA", "").replace("\\n", "\n").strip()
        if ca_inline:
            ca_archivo = os.path.join(tempfile.gettempdir(), "stock-aiven-ca.pem")
            with open(ca_archivo, "w", encoding="utf-8") as archivo:
                archivo.write(ca_inline + "\n")
            ca = ca_archivo
        else:
            ca = env.get("DB_SSL_CA_PATH", os.path.join("config", "ca.pem"))
            if not os.path.isabs(ca):
                ca = os.path.join(os.path.dirname(__file__), "..", ca)
            ca = os.path.abspath(ca)
        config["ssl_ca"] = ca
        config["ssl_verify_cert"] = True
    return mysql.connector.connect(**config)


def descargar(sheet_id, destino):
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx&t={int(time.time())}"
    pedido = urllib.request.Request(url, headers={"User-Agent": "stock-node-sync"})
    with urllib.request.urlopen(pedido, timeout=90) as respuesta:
        datos = respuesta.read()
    if not datos.startswith(b"PK"):
        raise PermissionError("respuesta_no_xlsx")
    with open(destino, "wb") as archivo:
        archivo.write(datos)
    return hashlib.sha256(datos).hexdigest()


def sincronizar_archivo(ruta):
    hojas = parsear_libro(ruta)
    cnx = conectar()
    try:
        aplicar(cnx, hojas)
    except Exception:
        cnx.rollback()
        raise
    finally:
        cnx.close()


def probar(ruta):
    hojas = parsear_libro(ruta)
    for hoja in hojas:
        print(f"== {hoja['pais']} movimientos={len(hoja['movimientos'])} categorias={len(hoja['categorias'])}")
        for categoria in hoja["categorias"]:
            print(f"  {categoria['stock']}\t{categoria['nombre']}")


def vigilar():
    env = leer_env()
    sheet_id = env.get("GOOGLE_SHEET_ID", "").strip()
    if not sheet_id:
        log.error("Falta GOOGLE_SHEET_ID en .env")
        sys.exit(1)
    intervalo = max(10, int(env.get("SYNC_INTERVAL_SEC", "15")))
    ultimo = None
    aviso = 0
    log.info("Sincronización de la planilla cada %s segundos", intervalo)
    while True:
        destino = os.path.join(tempfile.gettempdir(), "stock-planilla.xlsx")
        try:
            firma = descargar(sheet_id, destino)
            if firma != ultimo:
                sincronizar_archivo(destino)
                ultimo = firma
            aviso = 0
        except urllib.error.HTTPError as error:
            if error.code in (401, 403):
                ahora = time.time()
                if ahora - aviso > 60:
                    log.warning(AVISO_COMPARTIR)
                    aviso = ahora
            else:
                log.error("No se pudo descargar la planilla: HTTP %s", error.code)
        except Exception as error:
            log.error("La sincronización falló: %s", error)
        time.sleep(intervalo)


def una_vez():
    env = leer_env()
    sheet_id = env.get("GOOGLE_SHEET_ID", "").strip()
    if not sheet_id:
        log.error("Falta GOOGLE_SHEET_ID")
        sys.exit(1)
    destino = os.path.join(tempfile.gettempdir(), "stock-planilla.xlsx")
    descargar(sheet_id, destino)
    sincronizar_archivo(destino)


def main():
    if len(sys.argv) >= 2 and sys.argv[1] == "--una-vez":
        una_vez()
        return
    if len(sys.argv) >= 3 and sys.argv[1] == "--probar":
        probar(sys.argv[2])
        return
    if len(sys.argv) >= 3 and sys.argv[1] == "--archivo":
        sincronizar_archivo(sys.argv[2])
        return
    vigilar()


if __name__ == "__main__":
    main()
