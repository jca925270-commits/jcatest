"""La base de MAC es una sola hoja de cálculo. Este paso ya no copia otro archivo."""

import datetime as dt
import json
import os
import shutil
import subprocess
import sys
import tempfile
import urllib.request

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
from importar_stock_cantidad import limpiar_texto  # noqa: E402
from sync_planilla_drive import log  # noqa: E402
from sync_trazabilidad_mac import HOJAS, clasificar, mac, norm_header, sheet_id  # noqa: E402

EXCEL_ID = "1fNSvCBnYOqBd-IwC6lAdQw6cepZbsAIg8lHj2uGgloU"
PREFIJO = {tabla: prefijo for prefijo, _tipo, tabla in HOJAS}
TOPE_FILAS = 80
TOPE_CELDAS = 200


def excel_id():
    from sync_planilla_drive import leer_env
    return leer_env().get("GOOGLE_SHEET_MAC_EXCEL_ID", "").strip() or EXCEL_ID


def clasificar_espejo(nombre):
    limpio = (nombre or "").strip()
    if limpio.lower().startswith("engreso pantalla"):
        return "egreso", "trazabilidad_egreso_pantalla"
    return clasificar(nombre)


def visible(columna, valor):
    if isinstance(valor, dt.datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, dt.date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float) and valor == int(valor):
        valor = int(valor)
    if isinstance(valor, int) and not isinstance(valor, bool):
        texto = str(valor)
    else:
        texto = limpiar_texto(valor) or ""
    if "MAC" in columna:
        return mac(texto) or ""
    if texto.upper() in ("-", "#N/A", "#REF!", "#VALUE!"):
        return ""
    return texto


def descargar(sheet, destino):
    url = f"https://docs.google.com/spreadsheets/d/{sheet}/export?format=xlsx&t={int(dt.datetime.now().timestamp())}"
    pedido = urllib.request.Request(url, headers={"User-Agent": "stock-node-sync"})
    with urllib.request.urlopen(pedido, timeout=90) as respuesta:
        datos = respuesta.read()
    if not datos.startswith(b"PK"):
        raise PermissionError("respuesta_no_xlsx")
    with open(destino, "wb") as archivo:
        archivo.write(datos)


def leer(ruta):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hojas = {}
    try:
        for nombre in libro.sheetnames:
            clase = clasificar_espejo(nombre)
            if not clase:
                continue
            _tipo, tabla = clase
            if tabla in hojas:
                continue
            ws = libro[nombre]
            header = None
            cols = {}
            filas = {}
            for numero, row in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = list(row)
                if header is None and any(norm_header(v) == "MAC WIFI" for v in vals):
                    header = numero
                    for indice, valor in enumerate(vals):
                        clave = norm_header(valor)
                        if clave and clave not in cols:
                            cols[clave] = indice
                    continue
                if header is None or "MAC WIFI" not in cols:
                    continue
                indice_wifi = cols["MAC WIFI"]
                bruto = vals[indice_wifi] if indice_wifi < len(vals) else None
                wifi = mac(bruto)
                if not wifi or wifi in filas:
                    continue
                datos = {}
                for clave, indice in cols.items():
                    celda = vals[indice] if indice < len(vals) else None
                    datos[clave] = visible(clave, celda)
                filas[wifi] = datos
            hojas[tabla] = filas
    finally:
        libro.close()
    return hojas


def armar_plan(nativa, excel):
    copiar = []
    completar = []
    distintas = 0
    solo_nativa = 0
    for _prefijo, _tipo, tabla in HOJAS:
        filas_n = nativa.get(tabla, {})
        filas_e = excel.get(tabla, {})
        if tabla not in excel:
            continue
        for wifi, datos_n in filas_n.items():
            if wifi not in filas_e:
                solo_nativa += 1
        for wifi, datos_e in filas_e.items():
            if wifi not in filas_n:
                celdas = {clave: valor for clave, valor in datos_e.items() if valor}
                if celdas.get("MAC WIFI"):
                    copiar.append({"prefijo": PREFIJO[tabla], "celdas": celdas})
                continue
            datos_n = filas_n[wifi]
            for clave, valor_e in datos_e.items():
                valor_n = datos_n.get(clave, "")
                if valor_n == valor_e:
                    continue
                if not valor_n and valor_e:
                    completar.append({
                        "prefijo": PREFIJO[tabla],
                        "mac": wifi,
                        "columna": clave,
                        "valor": valor_e,
                    })
                elif valor_n and valor_e:
                    distintas += 1
    return copiar, completar, distintas, solo_nativa


def replicar_hacia_nativa(seco=False):
    log.info("MAC espejo: la base de MAC es una sola planilla y no se copia desde otro archivo.")
    return {"filas": 0, "celdas": 0}


if __name__ == "__main__":
    replicar_hacia_nativa(seco="--seco" in sys.argv)
