"""
importar_razon_social.py
--------------------------
Importa la lista completa de clientes (razón social) desde
razon_social.xlsx hacia la tabla `clientes`. Es una sola columna,
sin encabezado, un nombre por fila.

Es seguro correrlo de nuevo: no duplica (compara por nombre exacto,
igual que el resto del sistema).

Uso:
    python importar_razon_social.py "C:/ruta/al/razon_social.xlsx"
"""

import sys
import os
import logging

import mysql.connector
from openpyxl import load_workbook

from sync_watcher import Catalogos, DB_CFG

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("importar_razon_social")


def main():
    if len(sys.argv) < 2:
        print('Uso: python importar_razon_social.py "C:/ruta/al/razon_social.xlsx"')
        sys.exit(1)
    ruta = sys.argv[1]
    if not os.path.exists(ruta):
        log.error("No se encuentra el archivo: %s", ruta)
        sys.exit(1)

    wb = load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    cnx = mysql.connector.connect(**DB_CFG)
    cat = Catalogos(cnx)

    procesados = 0
    vistos = set()
    for row in ws.iter_rows(values_only=True):
        valor = row[0] if row else None
        if not valor:
            continue
        nombre = str(valor).strip()
        if not nombre or nombre in vistos:
            continue
        vistos.add(nombre)
        cat.cliente(nombre)
        procesados += 1

    cur = cnx.cursor()
    cur.execute("SELECT COUNT(*) FROM clientes")
    total_clientes = cur.fetchone()[0]
    cnx.close()

    log.info("Procesados %s nombres únicos de la planilla.", procesados)
    log.info("Total de clientes en la base ahora: %s", total_clientes)


if __name__ == "__main__":
    main()
