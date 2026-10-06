"""
Sincroniza Testeo EBD-Reparados, D3refresh y No pasaron prueba.

Las filas van a testeo_ebd_reparados, testeo_d3_refresh y
testeo_no_pasaron_prueba. No se copian a movimientos.
"""

import os
import re
import unicodedata
from datetime import datetime

from openpyxl import load_workbook

from sync_planilla_drive import conectar, descargar, leer_env, log
from sync_soporte_planilla import como_texto, fecha_o_texto

SHEET_TESTEO_ID = "16Ag2IJwMRHQSkrRIO8pC9ea4dxcSab6EC3TGKg0XmTI"

HOJAS = {
    "testeoebdreparados": ("ebd", "testeo_ebd_reparados", "Testeo EBD Reparados"),
    "d3refresh": ("d3", "testeo_d3_refresh", "D3 reutilizados/devueltos/sin antena wifi"),
    "nopasaronprueba": ("no_pasaron", "testeo_no_pasaron_prueba", "No pasaron prueba"),
}

# La pestaña No pasaron prueba no trae títulos: repite el orden de Testeo EBD.
POSICIONES = {
    "numero": 0,
    "fecha_comienzo": 2,
    "semanas": 3,
    "fecha_fin": 4,
    "mac": 5,
    "wifi": 6,
    "eth": 7,
    "tipo": 8,
    "observaciones": 9,
    "entidad": 10,
    "ubicacion": 11,
    "estado": 12,
}


def sheet_id():
    return leer_env().get("GOOGLE_SHEET_TESTEO_ID", "").strip() or SHEET_TESTEO_ID


def clave(valor):
    texto = (como_texto(valor, 200) or "").lower()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn"
    )
    texto = texto.replace(">", "")
    return re.sub(r"\s+", " ", texto).strip()


def indice(header, *nombres):
    claves = [clave(h) for h in header]
    for nombre in nombres:
        buscado = clave(nombre)
        for i, actual in enumerate(claves):
            if actual == buscado:
                return i
    for nombre in nombres:
        buscado = clave(nombre)
        for i, actual in enumerate(claves):
            if buscado and actual.startswith(buscado):
                return i
    return None


def celda(row, i):
    if i is None or i < 0 or i >= len(row):
        return None
    return row[i]


def tipo_hoja(nombre):
    limpio = re.sub(r"[\s\-]+", "", nombre).lower()
    return HOJAS.get(limpio)


def semanas(valor):
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float) and valor == int(valor):
        return str(int(valor))
    if isinstance(valor, int) and not isinstance(valor, bool):
        return str(valor)
    return como_texto(valor, 40)


def partir_macs(valor):
    texto = como_texto(valor, 80) or ""
    partes = re.findall(r"[0-9A-Fa-f]{2}(?::[0-9A-Fa-f]{2}){5}|[0-9A-Fa-f]{12}", texto)

    def una(parte):
        if not parte:
            return None
        limpio = re.sub(r"[^0-9A-Fa-f]", "", parte).upper()
        return limpio[:20] if len(limpio) >= 8 else None

    wifi = una(partes[0]) if partes else None
    eth = una(partes[1]) if len(partes) > 1 else None
    return wifi, eth


def observaciones_de(row, columnas, incluir_extras):
    principal = como_texto(celda(row, columnas.get("observaciones")), 4000)
    if not incluir_extras:
        return principal
    partes = []
    if principal:
        partes.append(principal)
    usados = {i for i in columnas.values() if i is not None}
    for i, valor in enumerate(row):
        if i in usados:
            continue
        texto = como_texto(valor, 4000)
        if not texto or texto == "." or texto.lower().startswith("total"):
            continue
        partes.append(texto)
    return " · ".join(partes) or None


def fila_de(row, numero_fila, columnas, con_titulo):
    mac_wifi, mac_eth = partir_macs(celda(row, columnas.get("mac")))
    if con_titulo:
        i_num = columnas.get("numero")
        bruto = celda(row, i_num)
        if bruto is None and i_num:
            bruto = celda(row, i_num - 1)
    else:
        bruto = celda(row, columnas.get("numero"))
    if isinstance(bruto, float) and bruto == int(bruto):
        numero = int(bruto)
    elif isinstance(bruto, int) and not isinstance(bruto, bool):
        numero = bruto
    else:
        numero = None
    if numero is None and not mac_wifi and not mac_eth:
        return None
    comienzo, comienzo_texto = fecha_o_texto(celda(row, columnas.get("fecha_comienzo")))
    fin, fin_texto = fecha_o_texto(celda(row, columnas.get("fecha_fin")))
    return {
        "numero": numero,
        "fecha_comienzo": comienzo,
        "fecha_comienzo_texto": comienzo_texto,
        "semanas": semanas(celda(row, columnas.get("semanas"))),
        "fecha_fin": fin,
        "fecha_fin_texto": fin_texto,
        "mac_wifi": mac_wifi,
        "mac_eth": mac_eth,
        "wifi": como_texto(celda(row, columnas.get("wifi")), 40),
        "eth": como_texto(celda(row, columnas.get("eth")), 40),
        "tipo": como_texto(celda(row, columnas.get("tipo")), 120),
        "observaciones": observaciones_de(row, columnas, not con_titulo),
        "entidad": como_texto(celda(row, columnas.get("entidad")), 255),
        "ubicacion": como_texto(celda(row, columnas.get("ubicacion")), 255),
        "estado": como_texto(celda(row, columnas.get("estado")), 80),
        "origen_fila": numero_fila,
    }


def columnas_de(header):
    return {
        "numero": indice(header, "numeracion"),
        "fecha_comienzo": indice(header, "fecha comienzo"),
        "semanas": indice(header, "semanas"),
        "fecha_fin": indice(header, "fecha finalizacion"),
        "mac": indice(header, "mac"),
        "wifi": indice(header, "wifi"),
        "eth": indice(header, "eth"),
        "tipo": indice(header, "tipo de fuente", "tipo de equipo"),
        "observaciones": indice(header, "observaciones"),
        "entidad": indice(header, "entidad"),
        "ubicacion": indice(header, "ubicacion", "orientacion"),
        "estado": indice(header, "estado"),
    }


def clave_total(valor):
    texto = clave(valor).replace(":", "").strip()
    if texto.startswith("total "):
        texto = texto[6:].strip()
    return texto[:40]


def totales_de(filas_crudas):
    vistos = set()
    salida = []
    for row in filas_crudas:
        for i, valor in enumerate(row):
            if not clave(valor).startswith("total"):
                continue
            bruto = row[i + 1] if i + 1 < len(row) else None
            if isinstance(bruto, float) and bruto == int(bruto):
                cantidad = int(bruto)
            elif isinstance(bruto, int) and not isinstance(bruto, bool):
                cantidad = bruto
            else:
                continue
            codigo = clave_total(valor)
            if not codigo or codigo in vistos:
                continue
            vistos.add(codigo)
            etiqueta = re.sub(r"\s+", " ", str(valor)).strip()
            salida.append({"clave": codigo, "etiqueta": etiqueta, "cantidad": cantidad})
    return salida


def momento(fila):
    return (fila["fecha_fin"] or fila["fecha_comienzo"] or "", fila["origen_fila"])


def quedar_una(filas):
    """Si la planilla repite la misma MAC, se conserva el registro más nuevo."""
    elegidas = {}
    orden = []
    sin_mac = []
    for fila in filas:
        if not fila["mac_wifi"] and not fila["mac_eth"]:
            sin_mac.append(fila)
            continue
        clave_mac = (fila["mac_wifi"], fila["mac_eth"])
        anterior = elegidas.get(clave_mac)
        if anterior is None:
            elegidas[clave_mac] = fila
            orden.append(clave_mac)
        elif momento(fila) >= momento(anterior):
            elegidas[clave_mac] = fila
    return sin_mac + [elegidas[clave_mac] for clave_mac in orden]


def parsear(ruta):
    libro = load_workbook(ruta, data_only=True, read_only=True)
    hojas = []
    try:
        for nombre in libro.sheetnames:
            info = tipo_hoja(nombre)
            if not info:
                continue
            tipo, tabla, titulo = info
            hoja = libro[nombre]
            filas_crudas = list(hoja.iter_rows(values_only=True))
            header = None
            inicio = 0
            for i, row in enumerate(filas_crudas):
                if indice(row, "mac") is not None and indice(row, "wifi") is not None:
                    header = list(row)
                    inicio = i + 1
                    break
            columnas = columnas_de(header) if header else POSICIONES
            filas = []
            for numero_fila, row in enumerate(filas_crudas[inicio:], start=inicio + 1):
                fila = fila_de(row, numero_fila, columnas, bool(header))
                if fila:
                    filas.append(fila)
            hojas.append({
                "nombre": nombre,
                "tipo": tipo,
                "tabla": tabla,
                "titulo": titulo,
                "filas": quedar_una(filas),
                "totales": totales_de(filas_crudas),
            })
    finally:
        libro.close()
    faltan = {info[0] for info in HOJAS.values()} - {h["tipo"] for h in hojas}
    if faltan:
        raise RuntimeError(f"Faltan pestañas de testeo: {', '.join(sorted(faltan))}")
    return hojas


def asegurar_tablas(cnx):
    ruta = os.path.join(os.path.dirname(__file__), "..", "sql", "migracion_testeo_planilla.sql")
    with open(ruta, encoding="utf-8") as archivo:
        script = archivo.read()
    cur = cnx.cursor()
    for sentencia in script.split(";"):
        sentencia = sentencia.strip()
        if sentencia:
            cur.execute(sentencia)
    cnx.commit()


def guardar_totales(cur, hoja, totales):
    if not totales:
        cur.execute("DELETE FROM testeo_totales WHERE hoja = %s", (hoja,))
        return
    cur.executemany(
        """INSERT INTO testeo_totales (hoja, clave, etiqueta, cantidad, orden)
           VALUES (%s, %s, %s, %s, %s)
           ON DUPLICATE KEY UPDATE
             etiqueta = VALUES(etiqueta), cantidad = VALUES(cantidad), orden = VALUES(orden)""",
        [(hoja, t["clave"], t["etiqueta"], t["cantidad"], i) for i, t in enumerate(totales)],
    )
    marcas = ",".join(["%s"] * len(totales))
    cur.execute(
        f"DELETE FROM testeo_totales WHERE hoja = %s AND clave NOT IN ({marcas})",
        [hoja] + [t["clave"] for t in totales],
    )


def aplicar(cnx, hojas):
    cur = cnx.cursor()
    total = 0
    for hoja in hojas:
        if not hoja["filas"]:
            log.warning("Testeo %s no trajo filas; no se borra la tabla", hoja["nombre"])
            continue
        tabla = hoja["tabla"]
        cur.executemany(
            f"""INSERT INTO {tabla}
                (numero, fecha_comienzo, fecha_comienzo_texto, semanas, fecha_fin, fecha_fin_texto,
                 mac_wifi, mac_eth, wifi, eth, tipo, observaciones, entidad, ubicacion, estado, origen_fila)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                  numero=VALUES(numero),
                  fecha_comienzo=VALUES(fecha_comienzo),
                  fecha_comienzo_texto=VALUES(fecha_comienzo_texto),
                  semanas=VALUES(semanas),
                  fecha_fin=VALUES(fecha_fin),
                  fecha_fin_texto=VALUES(fecha_fin_texto),
                  mac_wifi=VALUES(mac_wifi),
                  mac_eth=VALUES(mac_eth),
                  wifi=VALUES(wifi),
                  eth=VALUES(eth),
                  tipo=VALUES(tipo),
                  observaciones=VALUES(observaciones),
                  entidad=VALUES(entidad),
                  ubicacion=VALUES(ubicacion),
                  estado=VALUES(estado)""",
            [
                (f["numero"], f["fecha_comienzo"], f["fecha_comienzo_texto"], f["semanas"],
                 f["fecha_fin"], f["fecha_fin_texto"], f["mac_wifi"], f["mac_eth"], f["wifi"],
                 f["eth"], f["tipo"], f["observaciones"], f["entidad"], f["ubicacion"],
                 f["estado"], f["origen_fila"])
                for f in hoja["filas"]
            ],
        )
        cur.execute(
            """CREATE TEMPORARY TABLE IF NOT EXISTS tmp_testeo_sync (
                 origen_fila INT NOT NULL PRIMARY KEY
               )"""
        )
        cur.execute("DELETE FROM tmp_testeo_sync")
        cur.executemany(
            "INSERT IGNORE INTO tmp_testeo_sync (origen_fila) VALUES (%s)",
            [(f["origen_fila"],) for f in hoja["filas"]],
        )
        cur.execute(
            f"""DELETE t FROM {tabla} t
                LEFT JOIN tmp_testeo_sync s ON s.origen_fila = t.origen_fila
                WHERE s.origen_fila IS NULL"""
        )
        total += len(hoja["filas"])
        guardar_totales(cur, hoja["tipo"], hoja.get("totales") or [])
        log.info("Testeo %s: %s filas, %s totales", hoja["nombre"], len(hoja["filas"]), len(hoja.get("totales") or []))
    cur.execute(
        """INSERT INTO sync_control (archivo, ultima_sync, filas_ok, filas_error, detalle)
           VALUES ('TESTEO_PLANILLA', NOW(), %s, 0, 'Google Sheets')
           ON DUPLICATE KEY UPDATE
             ultima_sync = NOW(), filas_ok = VALUES(filas_ok),
             filas_error = 0, detalle = 'Google Sheets'""",
        (total,),
    )
    cnx.commit()
    log.info("Planilla de testeo actualizada: %s filas", total)


def aplicar_archivo(ruta):
    hojas = parsear(ruta)
    cnx = conectar()
    try:
        asegurar_tablas(cnx)
        aplicar(cnx, hojas)
    except Exception:
        cnx.rollback()
        raise
    finally:
        cnx.close()


def sincronizar():
    destino = os.path.join(os.environ.get("TEMP", os.environ.get("TMP", "/tmp")), "stock-testeo.xlsx")
    descargar(sheet_id(), destino)
    aplicar_archivo(destino)


if __name__ == "__main__":
    sincronizar()
